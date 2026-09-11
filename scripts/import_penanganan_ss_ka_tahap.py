# -*- coding: utf-8 -*-
"""Impor docs/New/11092026/.../[rencum asli] PENANGANAN SS KA PERTAHAP.xlsx
ke penanganan_ss_ka_tahap -- 3 sheet ("TAHAP 1 39 LOKASI", "TAHAP II 42
LOKASI", "TAHAP III 55 LOKASI") digabung + kolom tahap. Header row (dicari
dinamis per sheet lewat sel yg isinya "No" setelah di-strip -- posisi kolom
gesar-geser antar sheet, sheet Tahap 1 punya 3 sel kosong di depan vs 1 sel
di sheet lain) dipetakan by-name (bukan by-index tetap) supaya tahan
terhadap pergeseran itu.

Idempotent: DELETE + reinsert penuh (136 baris, tidak ada alasan utk upsert
per-baris).

Usage (venv aktif):
    python scripts/import_penanganan_ss_ka_tahap.py
"""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

import openpyxl

from db import db_cursor as pg_cursor  # noqa: E402

REPO_ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = REPO_ROOT / "docs" / "New" / "11092026" / "drive-download-20260911T021138Z-1-001" / "[rencum asli] PENANGANAN SS KA PERTAHAP.xlsx"

SHEETS = [
    ("TAHAP 1 39 LOKASI", "I"),
    ("TAHAP II 42 LOKASI", "II"),
    ("TAHAP III 55 LOKASI ", "III"),
]

COL_MAP = {
    "NO": "no",
    "NOTASI SS (KONSULTAN)": "notasi_ss_konsultan",
    "NOMOR NOTASI": "nomor_notasi",
    "NOTASI SS (KEPMEN367)": "notasi_ss_kepmen367",
    "NOTASI JPL": "notasi_jpl",
    "NAMA RUAS": "nama_ruas",
    "NOMOR RUAS": "nomor_ruas",
    "INDIKASI PANJANG PENANGANAN (M)": "indikasi_panjang_penanganan_m",
    "LOKASI JPL": "lokasi_jpl",
    "INDIKASI KEBUTUHAN FRONTAGE": "indikasi_kebutuhan_frontage",
    "PRAKIRAAN BIAYA KONSTRUKSI": "prakiraan_biaya_konstruksi",
    "KEBUTUHAN LAHAN (M2)": "kebutuhan_lahan_m2",
    "PRAKIRAAN BIAYA LAHAN": "prakiraan_biaya_lahan",
    "KETERSEDIAAN DESAIN": "ketersediaan_desain",
    "NILAI": "nilai",
    "TOTAL NILAI": "total_nilai",
    "RANGKING": "rangking",
    "EIRR": "eirr",
    "BCR": "bcr",
    "NPV": "npv",
}

INT_COLS = {"no", "nomor_notasi", "rangking"}
TEXT_COLS = {"notasi_ss_konsultan", "notasi_ss_kepmen367", "notasi_jpl", "nama_ruas", "nomor_ruas",
             "lokasi_jpl", "indikasi_kebutuhan_frontage", "ketersediaan_desain"}

TABLE_COLS = ["tahap"] + list(COL_MAP.values())


def _clean(v):
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def parse_sheet(ws, tahap):
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
    header_idx = next(i for i, r in enumerate(rows) if r and any(str(c).strip().upper() == "NO" for c in r if c is not None))
    header = [str(h).strip().upper() if h else None for h in rows[header_idx]]
    col_idx = {COL_MAP[h]: i for i, h in enumerate(header) if h in COL_MAP}

    out = []
    for r in rows[header_idx + 1:]:
        if not r or all(v is None for v in r):
            continue
        no_val = r[col_idx["no"]]
        if not isinstance(no_val, (int, float, str)) or (isinstance(no_val, str) and not no_val.strip()):
            continue
        rec = {"tahap": tahap}
        for col, idx in col_idx.items():
            val = _clean(r[idx]) if idx < len(r) else None
            if col in INT_COLS and val is not None:
                val = int(float(val))
            elif col in TEXT_COLS and isinstance(val, float):
                val = str(int(val)) if val.is_integer() else str(val)
            rec[col] = val
        out.append(rec)
    return out


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True, read_only=True)
    all_rows = []
    for sheet_name, tahap in SHEETS:
        rows = parse_sheet(wb[sheet_name], tahap)
        print(f"Tahap {tahap}: {len(rows)} baris")
        all_rows.extend(rows)

    records = [tuple(row.get(c) for c in TABLE_COLS) for row in all_rows]

    with pg_cursor() as cur:
        cur.execute(Path(__file__).with_name("schema_penanganan_ss_ka_tahap.sql").read_text(encoding="utf-8"))
        cur.execute("DELETE FROM penanganan_ss_ka_tahap")
        col_sql = ", ".join(TABLE_COLS)
        ph_sql = ", ".join(["%s"] * len(TABLE_COLS))
        cur.executemany(f"INSERT INTO penanganan_ss_ka_tahap ({col_sql}) VALUES ({ph_sql})", records)

    print(f"Selesai: {len(records)} baris diimpor ke penanganan_ss_ka_tahap.")


if __name__ == "__main__":
    main()
