"""Import "Paket LS Kegiatan Perpropinsi Persatker" Ditjen Bina Marga TA 2021-2025
(alokasi DIPA akhir APBN Bina Marga per paket x output) -> tabel paket_ls_bina_marga.

Sumber: docs/09102026/Paket LS/*.xlsx (Drive "Penambahan Data" no. 2, 28 Sep 2026),
satu file per tahun anggaran; TA dibaca dari judul sheet ("RENCANA PROGRAM DAN
KEGIATAN TA. 2024"), BUKAN dari nama file. Kajian: docs/kajian_paket_ls_sbsn.md.

Isinya sebagian besar JALAN NASIONAL (preservasi/pembangunan oleh satker PJN),
bukan IJD. Mulai TA 2023 ada RO "Dukungan Penanganan Jalan/Jembatan Daerah"
(kode RO .023/.024) = porsi IJD di DIPA Bina Marga, ditandai kolom
jalan_daerah. Ini ALOKASI anggaran, bukan realisasi. Tidak dipakai skor IJD/NPR.

Struktur sheet "Detail" (kolom penanda tingkat = kolom JUMLAH + 2):
  PROPINSI -> ringkasan PROGRAM/KEGIATAN/KRO/RO/KOMPONEN (dibaca utk nama kode)
           -> SATKER -> PROGRAM -> KEGIATAN -> LONGSEGMEN -> PAKET (per output)
                                            -> PAKET NON LS (per output)
Baris daun = PAKET + PAKET NON LS. Nilai di sumber dalam RIBUAN rupiah -> disimpan
dalam rupiah. Kolom sumber dana beda antar tahun (TA 2021: RPM/PDP/PHLN/SBSN;
2022+: RPM/RPLN/PHLN/LOCAL COST/SBSN) -> kolom tabel memuat gabungannya, NULL
bila tahun itu tak punya kolom tsb. Kode kegiatan berganti 2409 -> 7696 (2025).

Pengaman: impor DITOLAK bila jumlah daun per satker != baris SATKER, per provinsi
!= baris PROPINSI, atau total != baris "DITJEN BINA MARGA" di sheet Sumkeg
(sama dgn pola import_cer_awp1.py). NOMOR RUAS & LOKASI kosong di semua tahun:
lokasi hanya provinsi (+ nama ruas di dalam nama paket).

Wilayah: kode_provinsi BPS dicocokkan dari nama (wilayah_cocok.kunci_provinsi).
"PUSAT" -> NULL. TA 2021-2024 "PAPUA"/"PAPUA BARAT" = wilayah sebelum pemekaran
2022 (dicatat di catatan_data), dipetakan ke kode provinsi induknya.

DELETE + INSERT per tahun yang ada di folder, aman di-rerun.

Usage (venv aktif):
    python scripts/import_paket_ls_bina_marga.py --cek [folder]   # parse + rekonsiliasi saja, tanpa DB
    python scripts/import_paket_ls_bina_marga.py [folder]         # tulis ke DB (.env PG_*)
"""
import os
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
import psycopg
from psycopg.rows import dict_row

sys.path.insert(0, str(Path(__file__).resolve().parent))
from wilayah_cocok import kunci_provinsi  # noqa: E402

BASE_DIR = Path(__file__).resolve().parent.parent
DEFAULT_DIR = BASE_DIR / "docs" / "09102026" / "Paket LS"

# header sumber -> kolom tabel (nilai rupiah)
DANA = {"RPM": "rpm_rp", "PDP": "pdp_rp", "RPLN": "rpln_rp", "PHLN": "phln_rp",
        "LOCAL COST/RMP": "local_cost_rmp_rp", "SBSN": "sbsn_rp"}
KOLOM_DANA = list(DANA.values())
PAPUA_PRA_PEMEKARAN = {"PAPUA", "PAPUABARAT"}
# Porsi IJD di DIPA Bina Marga (RO .023 jalan / .024 jembatan, mulai TA 2023). SENGAJA tidak
# mencakup "Pembangunan/Preservasi Jalan Daerah Tertinggal dan Perbatasan" (RBC/RDC.006 dst.):
# itu jalan nasional di daerah tertinggal, bukan dukungan jalan daerah.
POLA_IJD = re.compile(r"Dukungan Penanganan (Jalan|Jembatan) Daerah", re.I)

DDL = """
CREATE TABLE IF NOT EXISTS paket_ls_bina_marga (
    id SERIAL PRIMARY KEY,
    tahun SMALLINT NOT NULL,
    provinsi TEXT,                  -- nama di sumber (huruf besar), 'PUSAT' utk satker pusat
    kode_provinsi SMALLINT,         -- ID BPS; NULL utk PUSAT
    satker_kode TEXT,
    satker_nama TEXT,
    kegiatan_kode TEXT,             -- 2409 (s.d. 2024) / 7696 (2025) dst.
    kegiatan_nama TEXT,
    jenis_paket TEXT,               -- 'LS' (bagian long segment) | 'NON LS'
    longsegmen_kode TEXT,
    longsegmen_nama TEXT,
    paket_kode TEXT,                -- no urut dlm long segment / kode huruf paket non LS
    paket_nama TEXT,
    kro_kode TEXT, kro_nama TEXT,   -- mis. CDC = O&M prasarana konektivitas darat (jalan)
    ro_kode TEXT, ro_nama TEXT,     -- mis. CDC.002 = Preservasi Rekonstruksi, Rehabilitasi Jalan.
                                    -- Nama NULL bila kosong di sumber (sebagian TA 2021); TIDAK dipinjam
                                    -- dari tahun lain: nomor RO bisa berganti arti antar tahun.
    komponen_kode TEXT,             -- KRO.RO.KOMPONEN, mis. CDC.002.321
    komponen_nama TEXT,             -- mis. Pemeliharaan Preventif
    jalan_daerah BOOLEAN,           -- RO "Dukungan Penanganan Jalan/Jembatan Daerah" (porsi IJD, TA 2023+)
    volume NUMERIC,
    satuan TEXT,                    -- Km / M / Dok / Unit ...
    rpm_rp NUMERIC, pdp_rp NUMERIC, rpln_rp NUMERIC, phln_rp NUMERIC,
    local_cost_rmp_rp NUMERIC, sbsn_rp NUMERIC,
    jumlah_rp NUMERIC,              -- total alokasi (rupiah; sumber dlm ribuan)
    keterangan TEXT,                -- kolom KETERANGAN sumber (mis. kode register SBSN 'M0330425,')
    catatan_data TEXT,
    sumber_file TEXT,
    diimpor_at TIMESTAMPTZ DEFAULT now()
);
CREATE INDEX IF NOT EXISTS idx_paket_ls_bm_tahun ON paket_ls_bina_marga (tahun);
CREATE INDEX IF NOT EXISTS idx_paket_ls_bm_prov ON paket_ls_bina_marga (kode_provinsi);
CREATE INDEX IF NOT EXISTS idx_paket_ls_bm_komp ON paket_ls_bina_marga (komponen_kode);
"""


def connect():
    from dotenv import load_dotenv
    load_dotenv(BASE_DIR / ".env")
    return psycopg.connect(
        host=os.environ.get("PG_HOST", "127.0.0.1"), port=int(os.environ.get("PG_PORT", "5432")),
        user=os.environ.get("PG_USER", "postgres"), password=os.environ.get("PG_PASS", ""),
        dbname=os.environ.get("PG_DB", "route_gis"), row_factory=dict_row,
    )


def _teks(v):
    s = re.sub(r"\s+", " ", str(v if v is not None else "")).strip()
    return s or None


def _angka(v):
    if v in (None, ""):
        return 0.0
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return 0.0


def _kode_output(s):
    """'2409.CDC.002.321 Pemeliharaan Preventif' / 'CDC.012.330 Pemeliharaan Rutin' /
    'CDC.023.409' -> ('CDC.002.321', 'Pemeliharaan Preventif' | None)."""
    m = re.match(r"^\s*(?:\d{4}\.)?([A-Z]{3}\.\d{3}\.\d{3})\s*(.*)$", str(s or ""))
    if not m:
        return None, _teks(s)
    return m.group(1), _teks(m.group(2))


def baca_file(path):
    """-> (tahun, daftar baris daun, ringkasan rekonsiliasi). Raise ValueError bila tak cocok."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Detail"]
    rows = list(ws.iter_rows(values_only=True))
    m = re.search(r"TA\.?\s*(\d{4})", str(rows[0][0] or ""))
    if not m:
        raise ValueError(f"{path.name}: tahun anggaran tak terbaca dari judul '{rows[0][0]}'")
    tahun = int(m.group(1))
    hdr = [str(x).strip().upper() if x else "" for x in rows[6]]
    ij = hdr.index("JUMLAH")
    iket, itag = ij + 1, ij + 2
    dana_idx = {i: DANA[h] for i, h in enumerate(hdr) if h in DANA}
    if len(dana_idx) != sum(1 for h in hdr[8:ij] if h):
        raise ValueError(f"{path.name}: kolom sumber dana tak dikenal: {hdr[8:ij]}")

    # Nama kode KRO/RO/KOMPONEN dari pohon ringkasan (kunci tanpa kode kegiatan)
    nama_kode = {}
    for r in rows[9:]:
        tag = _teks(r[itag]) if len(r) > itag else None
        if tag in ("KRO", "RO", "KOMPONEN") and r[0]:
            kode = re.sub(r"^\d{4}\.", "", str(r[0]).strip())
            nama_kode.setdefault(kode, _teks(r[1]))

    daun, prov, satker, keg, ls = [], None, None, (None, None), (None, None)
    sat_row, prov_row = {}, {}
    sum_sat, sum_prov = defaultdict(float), defaultdict(float)
    for r in rows[9:]:
        tag = (_teks(r[itag]) or "").upper() if len(r) > itag else ""
        if not tag:
            continue
        jumlah = _angka(r[ij])
        if tag == "PROPINSI":
            prov, satker = _teks(r[1]).upper(), None
            prov_row[prov] = jumlah
        elif tag == "SATKER":
            satker = (_teks(r[0]), _teks(r[1]))
            sat_row[(prov, satker[0])] = jumlah
        elif tag == "KEGIATAN" and satker:
            keg = (str(r[0]).strip().split(".")[-1], _teks(r[1]))
        elif tag == "LONGSEGMEN" and satker:
            ls = (_teks(r[0]), _teks(r[1]))
        elif tag in ("PAKET", "PAKET NON LS") and satker:
            kode, nama_out = _kode_output(r[2])
            kro = kode.split(".")[0] if kode else None
            ro = ".".join(kode.split(".")[:2]) if kode else None
            ro_nama = nama_kode.get(ro)
            dana = {k: None for k in KOLOM_DANA}
            for i, k in dana_idx.items():
                dana[k] = _angka(r[i]) * 1000
            daun.append({
                "tahun": tahun, "provinsi": prov, "satker_kode": satker[0], "satker_nama": satker[1],
                "kegiatan_kode": keg[0], "kegiatan_nama": keg[1],
                "jenis_paket": "LS" if tag == "PAKET" else "NON LS",
                "longsegmen_kode": ls[0] if tag == "PAKET" else None,
                "longsegmen_nama": ls[1] if tag == "PAKET" else None,
                "paket_kode": _teks(r[0]), "paket_nama": _teks(r[1]),
                "kro_kode": kro, "kro_nama": nama_kode.get(kro), "ro_kode": ro, "ro_nama": ro_nama,
                "komponen_kode": kode, "komponen_nama": nama_out or nama_kode.get(kode),
                "jalan_daerah": bool(ro_nama and POLA_IJD.search(ro_nama)),
                "volume": _angka(r[5]) if r[5] not in (None, "") else None, "satuan": _teks(r[7]),
                **dana, "jumlah_rp": jumlah * 1000, "keterangan": _teks(r[iket]) if len(r) > iket else None,
                "sumber_file": path.name,
            })
            sum_sat[(prov, satker[0])] += jumlah
            sum_prov[prov] += jumlah

    sk = list(wb["Sumkeg"].iter_rows(values_only=True))
    hk = [str(x).strip().upper() if x else "" for x in sk[6]]
    total_sumkeg = next(_angka(r[hk.index("JUMLAH")]) for r in sk[9:]
                        if r[1] and "DITJEN BINA MARGA" in str(r[1]).upper())
    salah = [f"satker {k}: baris {v:,.0f} vs paket {sum_sat[k]:,.0f}" for k, v in sat_row.items()
             if abs(v - sum_sat[k]) > 1]
    salah += [f"provinsi {k}: baris {v:,.0f} vs paket {sum_prov[k]:,.0f}" for k, v in prov_row.items()
              if abs(v - sum_prov[k]) > 1]
    total_daun = sum(sum_prov.values())
    if abs(total_daun - total_sumkeg) > 1:
        salah.append(f"total paket {total_daun:,.0f} vs Sumkeg {total_sumkeg:,.0f}")
    if salah:
        raise ValueError(f"{path.name} (TA {tahun}) TIDAK cocok dgn subtotal sumber:\n  " + "\n  ".join(salah[:20]))
    return tahun, daun, {"provinsi": len(prov_row), "satker": len(sat_row), "paket": len(daun),
                         "total_rp": total_daun * 1000}


def isi_wilayah(daun, prov_master):
    for d in daun:
        k = kunci_provinsi(d["provinsi"])
        d["kode_provinsi"] = prov_master.get(k)
        catatan = []
        if d["provinsi"] != "PUSAT" and d["kode_provinsi"] is None:
            catatan.append(f"provinsi '{d['provinsi']}' tak tercocokkan ke kode BPS")
        if k in PAPUA_PRA_PEMEKARAN and d["tahun"] <= 2024:
            catatan.append("wilayah provinsi sebelum pemekaran Papua 2022")
        if d["satuan"] in ("Km", "M") and not d["volume"] and d["jumlah_rp"]:
            catatan.append("volume 0/kosong (mis. pembayaran eskalasi) -- jangan dipakai utk biaya per satuan")
        d["catatan_data"] = "; ".join(catatan) or None


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    cek = "--cek" in sys.argv
    folder = Path(args[0]) if args else DEFAULT_DIR
    files = sorted(folder.glob("*.xlsx"))
    if not files:
        sys.exit(f"Tidak ada xlsx di {folder}")
    semua, per_tahun = [], {}
    for f in files:
        tahun, daun, ring = baca_file(f)  # ValueError -> berhenti, tidak ada yg ditulis
        if tahun in per_tahun:
            sys.exit(f"TA {tahun} muncul di dua file ({per_tahun[tahun]['file']} dan {f.name})")
        per_tahun[tahun] = {**ring, "file": f.name, "daerah": sum(d["jumlah_rp"] for d in daun if d["jalan_daerah"])}
        semua += daun
    for t, r in sorted(per_tahun.items()):
        print(f"TA {t}: {r['paket']:,} paket, {r['satker']} satker, {r['provinsi']} provinsi, "
              f"Rp {r['total_rp'] / 1e12:,.2f} T (jalan/jembatan daerah Rp {r['daerah'] / 1e12:,.2f} T) "
              f"-- cocok dgn subtotal sumber  [{r['file']}]")
    if cek:
        print("Mode --cek: tidak menulis ke database.")
        return
    conn = connect()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT DISTINCT kode_provinsi, provinsi FROM penduduk_kecamatan")
            prov_master = {kunci_provinsi(m["provinsi"]): int(m["kode_provinsi"]) for m in cur.fetchall()}
            isi_wilayah(semua, prov_master)
            cur.execute(DDL)
            cur.execute("DELETE FROM paket_ls_bina_marga WHERE tahun = ANY(%s)", (list(per_tahun),))
            kolom = ["tahun", "provinsi", "kode_provinsi", "satker_kode", "satker_nama", "kegiatan_kode",
                     "kegiatan_nama", "jenis_paket", "longsegmen_kode", "longsegmen_nama", "paket_kode",
                     "paket_nama", "kro_kode", "kro_nama", "ro_kode", "ro_nama", "komponen_kode",
                     "komponen_nama", "jalan_daerah", "volume", "satuan", *KOLOM_DANA, "jumlah_rp",
                     "keterangan", "catatan_data", "sumber_file"]
            with cur.copy(f"COPY paket_ls_bina_marga ({', '.join(kolom)}) FROM STDIN") as cp:
                for d in semua:
                    cp.write_row(tuple(d[k] for k in kolom))
        conn.commit()
        with conn.cursor() as cur:
            cur.execute("""SELECT tahun, COUNT(*) n, ROUND(SUM(jumlah_rp)/1e12, 2) t,
                                  COUNT(*) FILTER (WHERE kode_provinsi IS NULL AND provinsi <> 'PUSAT') tanpa_kode
                           FROM paket_ls_bina_marga GROUP BY 1 ORDER BY 1""")
            for r in cur.fetchall():
                print(f"DB {r['tahun']}: {r['n']:,} baris, Rp {r['t']} T, provinsi tanpa kode BPS {r['tanpa_kode']}")
    finally:
        conn.close()


if __name__ == "__main__":
    main()
