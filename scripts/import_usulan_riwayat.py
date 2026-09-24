"""Import riwayat usulan Inpres 2023-2026 (ekspor SITIA per tahun) ke
usulan_inpres_riwayat -- TERPISAH dari usulan_inpres.

Kenapa tabel terpisah: usulan_inpres diasumsikan berisi HANYA tarikan 2026 oleh
~30 query di app.py (ranking nasional, skor IJD, export, dashboard) dan puluhan
skrip pipeline, tak satu pun memfilter tahun_usulan. Memuat 2023-2025 ke tabel
itu akan mencampur semuanya secara diam-diam. Tabel ini slim (kolom yang
dibutuhkan untuk riwayat ruas & tren), kunci (tahun, id), aman di-rerun.
Kajian: docs/kajian_usulan_inpres_jalan_2023_2026.md.

Usage (venv aktif, .env berisi PG_*):
    python scripts/import_usulan_riwayat.py                  # semua *_usulan_inpres_*.xlsx di docs/24092026/Jalan
    python scripts/import_usulan_riwayat.py file1.xlsx ...   # file tertentu (tahun dari 4 digit awal nama file)
    python scripts/import_usulan_riwayat.py --tahun 2025 f.xlsx

Header antartahun berbeda (2023: "Nama Usulan", "Dukungan Surat DPR"); kolom yang
tidak ada di suatu tahun dibiarkan NULL, bukan 0. Nilai `Status Ruas` 2023-2024
(kode huruf K/P/O/N/S/NS) disimpan APA ADANYA -- artinya belum terdokumentasi.
"""
import argparse
import datetime
import glob
import os
import re
import sys
from pathlib import Path

import openpyxl
import psycopg

BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR))
DEFAULT_GLOB = str(BASE_DIR / "docs" / "24092026" / "Jalan" / "*_usulan_inpres_*.xlsx")

# alias header sumber (tahun lama) -> nama header kanonik
HEADER_ALIASES = {"Nama Usulan": "Nama Kegiatan"}

# header kanonik -> (kolom tabel, tipe: 's' teks, 'i' bigint, 'n' numeric, 'd' date)
COLUMNS = {
    "ID": ("id", "i"),
    "Nama Pengusul": ("nama_pengusul", "s"),
    "Provinsi": ("provinsi", "s"),
    "Kabupaten/Kota": ("kabupaten_kota", "s"),
    "Nama Kegiatan": ("nama_kegiatan", "s"),
    "Prioritas": ("prioritas", "i"),
    "Tgl. Pengusulan": ("tgl_pengusulan", "d"),
    "Kode Ruas": ("kode_ruas", "s"),
    "Nama Ruas": ("nama_ruas", "s"),
    "Panjang Ruas (KM)": ("panjang_ruas_km", "n"),
    "Status Ruas": ("status_ruas", "s"),
    "Jenis Penanganan": ("jenis_penanganan", "s"),
    "Panjang Penanganan (Pemda)": ("panjang_penanganan_pemda", "n"),
    "Alokasi Usulan (Pemda)": ("alokasi_usulan_pemda", "n"),
    "Alokasi Usulan (Balai)": ("alokasi_usulan_balai", "n"),
    "Alokasi Usulan (Kompetensi)": ("alokasi_usulan_kompetensi", "n"),
    "Tematik Kawasan (Kompetensi)": ("tematik_kawasan_kompetensi", "s"),
    "Seleksi Sistem": ("seleksi_sistem", "s"),
    "Kelengkapan Input Data": ("kelengkapan_input_data", "s"),
    "Verifikasi Balai": ("verifikasi_balai", "s"),
    "Verifikasi Kompetensi": ("verifikasi_kompetensi", "s"),
    "Verifikasi PFID": ("verifikasi_pfid", "s"),
    "Bilateral Bappenas": ("bilateral_bappenas", "s"),
    "Bilateral PU": ("bilateral_pu", "s"),
    "Alokasi/Nilai RDPP": ("nilai_rdpp", "n"),
    "Alokasi/Nilai DPP": ("nilai_dpp", "n"),
    "Diprogramkan": ("diprogramkan", "s"),
    "Penuntasan IJD Sebelumnya (Kompetensi)": ("penuntasan_ijd_kompetensi", "s"),
}

DDL = """
CREATE TABLE IF NOT EXISTS usulan_inpres_riwayat (
    tahun SMALLINT NOT NULL,
    id BIGINT NOT NULL,
    nama_pengusul TEXT, provinsi TEXT, kabupaten_kota TEXT, nama_kegiatan TEXT,
    prioritas INTEGER, tgl_pengusulan DATE,
    kode_ruas TEXT, nama_ruas TEXT, panjang_ruas_km NUMERIC, status_ruas TEXT,
    jenis_penanganan TEXT, panjang_penanganan_pemda NUMERIC,
    alokasi_usulan_pemda NUMERIC, alokasi_usulan_balai NUMERIC, alokasi_usulan_kompetensi NUMERIC,
    tematik_kawasan_kompetensi TEXT,
    seleksi_sistem TEXT, kelengkapan_input_data TEXT,
    verifikasi_balai TEXT, verifikasi_kompetensi TEXT, verifikasi_pfid TEXT,
    bilateral_bappenas TEXT, bilateral_pu TEXT,
    nilai_rdpp NUMERIC, nilai_dpp NUMERIC, diprogramkan TEXT,
    penuntasan_ijd_kompetensi TEXT,
    kode_provinsi SMALLINT, kode_kabupaten INTEGER,  -- ID BPS, diisi wilayah_id.isi_kode_riwayat
    sumber_file TEXT, diimpor_at TIMESTAMPTZ DEFAULT now(),
    PRIMARY KEY (tahun, id)
);
CREATE INDEX IF NOT EXISTS idx_usulan_riwayat_kode_ruas ON usulan_inpres_riwayat (kode_ruas);
"""


def clean(value, tipe):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    try:
        if tipe == "i":
            return int(float(value))
        if tipe == "n":
            return float(value)
        if tipe == "d":
            if isinstance(value, datetime.datetime):
                return value.date()
            if isinstance(value, datetime.date):
                return value
            return datetime.datetime.strptime(str(value)[:10], "%Y-%m-%d").date()
    except (ValueError, TypeError):
        return None
    return str(value).strip()


def tahun_dari_nama(path):
    m = re.match(r"(\d{4})_", os.path.basename(path))
    return int(m.group(1)) if m else None


def import_file(conn, path, tahun):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Worksheet"] if "Worksheet" in wb.sheetnames else wb.active
    rows = ws.iter_rows(values_only=True)
    header = [HEADER_ALIASES.get(h, h) for h in next(rows)]
    idx = {name: i for i, name in enumerate(header)}
    if "ID" not in idx:
        raise ValueError(f"{path}: kolom ID tidak ada")
    aktif = {h: c for h, c in COLUMNS.items() if h in idx}
    absen = [h for h in COLUMNS if h not in idx]
    kolom = ["tahun"] + [c[0] for c in aktif.values()] + ["sumber_file"]
    sql = (f"INSERT INTO usulan_inpres_riwayat ({', '.join(kolom)}) VALUES ({', '.join(['%s'] * len(kolom))}) "
           f"ON CONFLICT (tahun, id) DO UPDATE SET " +
           ", ".join(f"{k}=EXCLUDED.{k}" for k in kolom if k not in ("tahun", "id")) + ", diimpor_at=now()")
    batch, n = [], 0
    with conn.cursor() as cur:
        for row in rows:
            row = tuple(row) + (None,) * (len(header) - len(row))
            rec = [tahun]
            for h, (_c, tipe) in aktif.items():
                rec.append(clean(row[idx[h]], tipe))
            if rec[1 + list(aktif).index("ID")] is None:
                continue
            rec.append(os.path.basename(path))
            batch.append(tuple(rec))
            n += 1
            if len(batch) >= 500:
                cur.executemany(sql, batch)
                batch.clear()
        if batch:
            cur.executemany(sql, batch)
    conn.commit()
    return n, absen


def main():
    ap = argparse.ArgumentParser(description="Import riwayat usulan Inpres 2023-2026 -> usulan_inpres_riwayat")
    ap.add_argument("xlsx", nargs="*", help="file xlsx (default: semua di docs/24092026/Jalan)")
    ap.add_argument("--tahun", type=int, help="paksa tahun (default: 4 digit awal nama file)")
    args = ap.parse_args()

    from dotenv import load_dotenv
    load_dotenv(BASE_DIR / ".env")
    files = args.xlsx or sorted(glob.glob(DEFAULT_GLOB))
    if not files:
        sys.exit("Tidak ada file xlsx.")
    conn = psycopg.connect(
        host=os.environ.get("PG_HOST", "127.0.0.1"), port=int(os.environ.get("PG_PORT", "5432")),
        user=os.environ.get("PG_USER", "postgres"), password=os.environ.get("PG_PASS", ""),
        dbname=os.environ.get("PG_DB", "route_gis"))
    try:
        with conn.cursor() as cur:
            cur.execute(DDL)
        conn.commit()
        for f in files:
            tahun = args.tahun or tahun_dari_nama(f)
            if not tahun:
                sys.exit(f"Tahun tidak bisa ditebak dari nama file {f}; pakai --tahun")
            n, absen = import_file(conn, f, tahun)
            print(f"{tahun}: {n} baris dari {os.path.basename(f)} ({len(absen)} kolom tak ada di file -> NULL)")
        import wilayah_id  # ID wilayah: ref_wilayah dulu, lalu kode_provinsi/kode_kabupaten riwayat
        with conn.cursor() as cur:
            wilayah_id.build_ref_wilayah(cur)
            print("Kode wilayah riwayat terisi:", wilayah_id.isi_kode_riwayat(cur))
        conn.commit()
        with conn.cursor() as cur:
            cur.execute("SELECT tahun, COUNT(*) FROM usulan_inpres_riwayat GROUP BY 1 ORDER BY 1")
            print("Total di tabel:", cur.fetchall())
    finally:
        conn.close()


if __name__ == "__main__":
    main()
