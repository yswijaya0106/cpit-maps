"""Penyimpanan hasil analisis chat ("dataset") + export (xlsx/csv/geojson) +
laporan Word dari markdown.

Dipakai chat_providers.py (tool jalankan_query_sql/tampilkan_tabel/
buat_grafik/tampilkan_di_peta/buat_laporan) dan route /api/chat/* di app.py.

Kenapa disimpan di disk (.cache/chat_dataset), bukan dict in-memory: staging
menjalankan uvicorn --workers 2 -- request unduh bisa jatuh ke worker lain
dari yang menjalankan query. Direktori .cache/ sudah dipakai cache layer peta
dgn alasan yg sama. Berkas > MAKS_UMUR_HARI dihapus otomatis saat menyimpan.

Model LLM hanya menerima PRATINJAU hasil (beberapa puluh baris + jumlah
total); tabel/grafik/peta/export membaca dataset lengkap langsung dari sini,
jadi batas baris model tidak memotong hasil yang diunduh pengguna.
"""
import csv
import gzip
import io
import json
import re
import time
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from fastapi.encoders import jsonable_encoder

ROOT = Path(__file__).resolve().parent
DIR_DATASET = ROOT / ".cache" / "chat_dataset"
DIR_FILE = ROOT / ".cache" / "chat_file"
MAKS_BARIS = 20000
MAKS_UMUR_HARI = 7
_ID_RE = re.compile(r"^[a-f0-9]{16}$")


def _id_valid(ds_id: str) -> bool:
    return bool(ds_id) and bool(_ID_RE.match(ds_id))


def _bersihkan_lama(direktori: Path):
    batas = time.time() - MAKS_UMUR_HARI * 86400
    for p in direktori.glob("*"):
        try:
            if p.stat().st_mtime < batas:
                p.unlink()
        except OSError:
            pass


def simpan(columns: List[str], rows: List[list], judul: str = "", sql: str = "",
           pengguna: Optional[str] = None, terpotong: bool = False) -> str:
    DIR_DATASET.mkdir(parents=True, exist_ok=True)
    _bersihkan_lama(DIR_DATASET)
    ds_id = uuid.uuid4().hex[:16]
    isi = {
        "id": ds_id, "judul": judul or "Hasil query", "sql": sql, "pengguna": pengguna,
        "dibuat": datetime.now().isoformat(timespec="seconds"),
        "columns": columns, "rows": jsonable_encoder(rows), "terpotong": terpotong,
    }
    with gzip.open(DIR_DATASET / f"{ds_id}.json.gz", "wt", encoding="utf-8") as f:
        json.dump(isi, f, ensure_ascii=False)
    return ds_id


def muat(ds_id: str) -> Optional[Dict[str, Any]]:
    if not _id_valid(ds_id):
        return None
    p = DIR_DATASET / f"{ds_id}.json.gz"
    if not p.exists():
        return None
    with gzip.open(p, "rt", encoding="utf-8") as f:
        return json.load(f)


def ringkas(ds: Dict[str, Any]) -> Dict[str, Any]:
    return {"dataset_id": ds["id"], "judul": ds["judul"], "jumlah_baris": len(ds["rows"]),
            "kolom": ds["columns"], "terpotong": ds.get("terpotong", False)}


# ---------------------------------------------------------------------------
# Geometri: kolom GeoJSON (hasil ST_AsGeoJSON) atau pasangan lat/lon
# ---------------------------------------------------------------------------

_NAMA_LAT = ("lat", "latitude", "lintang", "y", "koordinat_lat")
_NAMA_LON = ("lon", "lng", "long", "longitude", "bujur", "x", "koordinat_lon")


def _sebagai_geojson(v):
    if isinstance(v, dict) and "type" in v and ("coordinates" in v or "geometries" in v):
        return v
    if isinstance(v, str) and v.lstrip().startswith("{") and '"type"' in v:
        try:
            g = json.loads(v)
            return g if isinstance(g, dict) and "type" in g else None
        except ValueError:
            return None
    return None


def deteksi_geometri(ds: Dict[str, Any]) -> Dict[str, Optional[str]]:
    """Tebak kolom geometri: kolom berisi GeoJSON, atau pasangan lat/lon."""
    cols = ds["columns"]
    sampel = ds["rows"][:20]
    for i, c in enumerate(cols):
        if any(_sebagai_geojson(r[i]) for r in sampel):
            return {"kolom_geometri": c, "kolom_lat": None, "kolom_lon": None}
    low = {c.lower(): c for c in cols}
    lat = next((low[n] for n in _NAMA_LAT if n in low), None)
    lon = next((low[n] for n in _NAMA_LON if n in low), None)
    if lat and lon:
        return {"kolom_geometri": None, "kolom_lat": lat, "kolom_lon": lon}
    return {"kolom_geometri": None, "kolom_lat": None, "kolom_lon": None}


def ke_geojson(ds: Dict[str, Any], kolom_geometri=None, kolom_lat=None, kolom_lon=None) -> Dict[str, Any]:
    if not (kolom_geometri or (kolom_lat and kolom_lon)):
        auto = deteksi_geometri(ds)
        kolom_geometri, kolom_lat, kolom_lon = auto["kolom_geometri"], auto["kolom_lat"], auto["kolom_lon"]
    cols = ds["columns"]
    idx = {c: i for i, c in enumerate(cols)}
    fitur = []
    for r in ds["rows"]:
        geom = None
        if kolom_geometri and kolom_geometri in idx:
            geom = _sebagai_geojson(r[idx[kolom_geometri]])
        elif kolom_lat in idx and kolom_lon in idx:
            try:
                la, lo = float(r[idx[kolom_lat]]), float(r[idx[kolom_lon]])
                if -90 <= la <= 90 and -180 <= lo <= 180 and not (la == 0 and lo == 0):
                    geom = {"type": "Point", "coordinates": [lo, la]}
            except (TypeError, ValueError):
                geom = None
        if not geom:
            continue
        props = {c: r[i] for c, i in idx.items() if c != kolom_geometri and not _sebagai_geojson(r[i])}
        fitur.append({"type": "Feature", "geometry": geom, "properties": props})
    return {"type": "FeatureCollection", "features": fitur,
            "kolom_geometri": kolom_geometri, "kolom_lat": kolom_lat, "kolom_lon": kolom_lon}


# ---------------------------------------------------------------------------
# Export tabel
# ---------------------------------------------------------------------------

def _nama_berkas(judul: str, ext: str) -> str:
    dasar = re.sub(r"[^A-Za-z0-9]+", "_", judul or "hasil").strip("_")[:60] or "hasil"
    return f"{dasar}_{datetime.now():%Y%m%d_%H%M}.{ext}"


def _kolom_tampil(ds):
    """Kolom GeoJSON mentah tidak ikut ke tabel xlsx/csv (terlalu panjang, tak terbaca)."""
    sampel = ds["rows"][:20]
    return [i for i, c in enumerate(ds["columns"]) if not any(_sebagai_geojson(r[i]) for r in sampel)]


def ekspor_csv(ds) -> tuple:
    keep = _kolom_tampil(ds)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([ds["columns"][i] for i in keep])
    for r in ds["rows"]:
        w.writerow(["" if r[i] is None else r[i] for i in keep])
    return ("﻿" + buf.getvalue()).encode("utf-8"), "text/csv", _nama_berkas(ds["judul"], "csv")


def ekspor_xlsx(ds) -> tuple:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    keep = _kolom_tampil(ds)
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    kepala = [ds["columns"][i] for i in keep]
    ws.append(kepala)
    for r in ds["rows"]:
        ws.append([r[i] if not isinstance(r[i], (dict, list)) else json.dumps(r[i], ensure_ascii=False) for i in keep])
    isi = PatternFill("solid", fgColor="0F2A4A")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = isi
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for j, k in enumerate(kepala, start=1):
        panjang = max([len(str(k))] + [len(str(r[keep[j - 1]] or "")) for r in ds["rows"][:200]])
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, panjang + 2), 60)

    meta = wb.create_sheet("Metode")
    for baris in (
        ("Judul", ds["judul"]), ("Dibuat", ds.get("dibuat", "")), ("Jumlah baris", len(ds["rows"])),
        ("Terpotong", f"Ya (maks. {MAKS_BARIS} baris)" if ds.get("terpotong") else "Tidak"),
        ("Sumber", "Asisten The Next - SiJalan — query read-only ke database aplikasi"),
        ("Query SQL", ds.get("sql", "")),
        ("Catatan", "Draf analisis AI — verifikasi sebelum dipakai resmi."),
    ):
        meta.append(list(baris))
    meta.column_dimensions["A"].width = 16
    meta.column_dimensions["B"].width = 110
    for row in meta.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    out = io.BytesIO()
    wb.save(out)
    return (out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            _nama_berkas(ds["judul"], "xlsx"))


def ekspor_geojson(ds, **kw) -> tuple:
    fc = ke_geojson(ds, **kw)
    body = {"type": "FeatureCollection", "name": ds["judul"], "features": fc["features"]}
    return (json.dumps(jsonable_encoder(body), ensure_ascii=False).encode("utf-8"), "application/geo+json",
            _nama_berkas(ds["judul"], "geojson"))


# ---------------------------------------------------------------------------
# Laporan Word dari markdown + lampiran dataset
# ---------------------------------------------------------------------------

_NAVY = (15, 42, 74)


def _tulis_inline(par, teks: str, ukuran=None):
    """**tebal**, *miring*, `kode` -> run python-docx."""
    from docx.shared import Pt

    for bagian in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)", teks):
        if not bagian:
            continue
        if bagian.startswith("**") and bagian.endswith("**"):
            run = par.add_run(bagian[2:-2])
            run.bold = True
        elif bagian.startswith("`") and bagian.endswith("`"):
            run = par.add_run(bagian[1:-1])
            run.font.name = "Consolas"
        elif bagian.startswith("*") and bagian.endswith("*") and len(bagian) > 2:
            run = par.add_run(bagian[1:-1])
            run.italic = True
        else:
            run = par.add_run(bagian)
        if ukuran:
            run.font.size = Pt(ukuran)


def _tabel_docx(doc, kepala: List[str], baris: List[List[str]], ukuran=8):
    from docx.shared import Pt, RGBColor
    from map_print import _docx_repeat_header, _docx_shading

    tb = doc.add_table(rows=1, cols=len(kepala))
    tb.style = "Table Grid"
    _docx_repeat_header(tb.rows[0])
    for i, k in enumerate(kepala):
        sel = tb.rows[0].cells[i]
        _docx_shading(sel, "%02X%02X%02X" % _NAVY)
        run = sel.paragraphs[0].add_run(str(k))
        run.bold = True
        run.font.size = Pt(ukuran)
        run.font.color.rgb = RGBColor(255, 255, 255)
    for ri, r in enumerate(baris):
        sel = tb.add_row().cells
        for i, v in enumerate(r[: len(kepala)]):
            _tulis_inline(sel[i].paragraphs[0], "" if v is None else str(v), ukuran)
            if ri % 2 == 1:
                _docx_shading(sel[i], "F3F6FB")
    return tb


def _markdown_ke_docx(doc, md: str):
    from docx.shared import Pt, RGBColor

    baris = md.replace("\r\n", "\n").split("\n")
    i = 0
    while i < len(baris):
        s = baris[i].rstrip()
        if not s.strip():
            i += 1
            continue
        # tabel markdown: | a | b | + baris pemisah |---|
        if s.lstrip().startswith("|") and i + 1 < len(baris) and re.match(r"^\s*\|?\s*:?-{2,}", baris[i + 1]):
            sel = lambda t: [x.strip() for x in t.strip().strip("|").split("|")]
            kepala = sel(s)
            isi = []
            i += 2
            while i < len(baris) and baris[i].lstrip().startswith("|"):
                isi.append(sel(baris[i]))
                i += 1
            _tabel_docx(doc, kepala, isi, 8.5)
            doc.add_paragraph()
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            h = doc.add_heading(level=min(len(m.group(1)), 3))
            run = h.add_run(m.group(2).replace("**", ""))
            run.font.color.rgb = RGBColor(*_NAVY)
            i += 1
            continue
        m = re.match(r"^\s*[-*]\s+(.*)", s)
        if m:
            _tulis_inline(doc.add_paragraph(style="List Bullet"), m.group(1))
            i += 1
            continue
        m = re.match(r"^\s*\d+[.)]\s+(.*)", s)
        if m:
            _tulis_inline(doc.add_paragraph(style="List Number"), m.group(1))
            i += 1
            continue
        if s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(12)
            _tulis_inline(p, s.lstrip("> "))
            i += 1
            continue
        _tulis_inline(doc.add_paragraph(), s)
        i += 1


def buat_laporan_docx(judul: str, isi_markdown: str, dataset_ids: List[str], pengguna: Optional[str]) -> Dict[str, Any]:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, Pt, RGBColor
    from map_print import _docx_cell_margin, _docx_field, _docx_shading, _tanggal_id

    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, attr, Mm(20))
    sekarang = datetime.now()

    pita = doc.add_table(rows=1, cols=1)
    sel = pita.rows[0].cells[0]
    _docx_shading(sel, "%02X%02X%02X" % _NAVY)
    r = sel.paragraphs[0].add_run(judul or "Laporan Analisis")
    r.bold = True
    r.font.size = Pt(17)
    r.font.color.rgb = RGBColor(255, 255, 255)
    r2 = sel.add_paragraph().add_run(f"THE NEXT - SiJalan  ·  {_tanggal_id(sekarang)}" + (f"  ·  {pengguna}" if pengguna else ""))
    r2.bold = True
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = RGBColor(34, 211, 165)
    _docx_cell_margin(pita, 3)

    cat = doc.add_paragraph()
    cat.paragraph_format.space_before = Pt(6)
    rc = cat.add_run("Draf analisis yang disusun asisten AI dari database aplikasi — verifikasi angka sebelum dipakai resmi.")
    rc.italic = True
    rc.font.size = Pt(8.5)
    rc.font.color.rgb = RGBColor(100, 110, 125)

    _markdown_ke_docx(doc, isi_markdown or "")

    lampiran = [d for d in (muat(x) for x in (dataset_ids or [])) if d]
    for n, ds in enumerate(lampiran, start=1):
        doc.add_page_break() if n == 1 else doc.add_paragraph()
        h = doc.add_heading(level=2)
        hr = h.add_run(f"Lampiran {n}. {ds['judul']}")
        hr.font.color.rgb = RGBColor(*_NAVY)
        keep = _kolom_tampil(ds)[:10]
        maks = 300
        ket = doc.add_paragraph()
        kr = ket.add_run(f"{len(ds['rows']):,} baris".replace(",", ".")
                         + (f" — ditampilkan {maks} pertama; data lengkap tersedia di unduhan Excel" if len(ds["rows"]) > maks else "")
                         + (f"; {len(ds['columns']) - len(keep)} kolom lain tidak ditampilkan" if len(ds["columns"]) > len(keep) else ""))
        kr.font.size = Pt(8)
        kr.font.color.rgb = RGBColor(90, 100, 115)
        _tabel_docx(doc, [ds["columns"][i] for i in keep],
                    [[ds_row[i] for i in keep] for ds_row in ds["rows"][:maks]], 7.5)
        if ds.get("sql"):
            q = doc.add_paragraph()
            q.paragraph_format.space_before = Pt(4)
            qr = q.add_run("Query: " + ds["sql"])
            qr.font.size = Pt(7)
            qr.font.name = "Consolas"
            qr.font.color.rgb = RGBColor(100, 110, 125)

    pk = sec.footer.paragraphs[0]
    pk.alignment = WD_ALIGN_PARAGRAPH.LEFT
    rk = pk.add_run(f"The Next - SiJalan · {_tanggal_id(sekarang)}    |    Halaman ")
    rk.font.size = Pt(7.5)
    rk2 = pk.add_run()
    rk2.font.size = Pt(7.5)
    _docx_field(rk2, "PAGE")

    DIR_FILE.mkdir(parents=True, exist_ok=True)
    _bersihkan_lama(DIR_FILE)
    file_id = uuid.uuid4().hex[:16]
    nama = _nama_berkas(judul or "laporan", "docx")
    buf = io.BytesIO()
    doc.save(buf)
    (DIR_FILE / f"{file_id}.docx").write_bytes(buf.getvalue())
    (DIR_FILE / f"{file_id}.json").write_text(json.dumps({"nama": nama, "judul": judul}), encoding="utf-8")
    return {"file_id": file_id, "nama_berkas": nama, "lampiran": len(lampiran)}


def muat_file(file_id: str):
    if not _id_valid(file_id):
        return None
    p, m = DIR_FILE / f"{file_id}.docx", DIR_FILE / f"{file_id}.json"
    if not p.exists():
        return None
    nama = json.loads(m.read_text(encoding="utf-8")).get("nama", "laporan.docx") if m.exists() else "laporan.docx"
    return p.read_bytes(), nama
