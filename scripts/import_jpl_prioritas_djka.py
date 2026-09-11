# -*- coding: utf-8 -*-
"""Impor docs/New/11092026/.../(REV) JPL_Prioritas_DJKA_v3.xlsx ke
jpl_prioritas_djka -- 7 sheet regional (Jakarta/Bandung/Semarang/Surabaya/
Medan/Padang/Palembang) digabung, header row (baris ke-3 tiap sheet) dibaca
dinamis per sheet karena 2 sheet (Bandung, Surabaya) punya kolom tambahan
"JUSTIFIKASI" dan 1 sheet (Palembang) punya "JML KECELAKAAN" yang tidak ada
di sheet lain. Sheet "RINGKASAN"/"Sheet3"/"MASTER"/"Sheet1" dilewati --
RINGKASAN & Sheet3 adalah rekap pivot (agregat dari data yg sama), MASTER &
Sheet1 subset/duplikat sheet regional (diverifikasi kolom "PETAK STASIUN"
baris pertamanya identik dgn sheet JAKARTA).

Idempotent: DELETE + reinsert penuh (~86 baris, tidak ada alasan utk upsert
per-baris).

Usage (venv aktif):
    python scripts/import_jpl_prioritas_djka.py
"""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

import openpyxl

from db import db_cursor as pg_cursor  # noqa: E402

REPO_ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = REPO_ROOT / "docs" / "New" / "11092026" / "drive-download-20260911T021138Z-1-001" / "(REV) JPL_Prioritas_DJKA_v3.xlsx"

REGIONAL_SHEETS = ["JAKARTA", "BANDUNG", "SEMARANG", "SURABAYA", "MEDAN", "PADANG", "PALEMBANG"]

# Header sumber -> kolom tabel.
COL_MAP = {
    "NO": "no",
    "PETAK STASIUN": "petak_stasiun",
    "NO JPL": "no_jpl",
    "LOKASI KM+HM": "lokasi_km_hm",
    "KOTA / KAB": "kota_kab",
    "STATUS PENJAGAAN": "status_penjagaan",
    "PROVINSI": "provinsi",
    "KATEGORI JALAN": "kategori_jalan",
    "LEBAR JALAN (m)": "lebar_jalan_m",
    "JENIS JALUR KA": "jenis_jalur_ka",
    "FREKUENSI KA": "frekuensi_ka",
    "HEADWAY KA": "headway_ka",
    "NAMA JALAN": "nama_jalan",
    "DAOP / DIVRE": "daop_divre",
    "JUSTIFIKASI": "justifikasi",
    "JML KECELAKAAN": "jumlah_kecelakaan",
}

TABLE_COLS = [
    "btp_wilayah_kerja", "no", "petak_stasiun", "no_jpl", "lokasi_km_hm", "kota_kab",
    "status_penjagaan", "provinsi", "kategori_jalan", "lebar_jalan_m", "jenis_jalur_ka",
    "frekuensi_ka", "headway_ka", "nama_jalan", "daop_divre", "justifikasi", "jumlah_kecelakaan",
]


def _clean(v):
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def parse_sheet(ws, wilayah):
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
    header_idx = next(i for i, r in enumerate(rows) if r and r[0] == "NO")
    header = [str(h).strip() if h else None for h in rows[header_idx]]
    col_idx = {COL_MAP[h]: i for i, h in enumerate(header) if h in COL_MAP}

    out = []
    for r in rows[header_idx + 1:]:
        if not r or all(v is None for v in r):
            continue
        no_val = r[col_idx["no"]]
        if not isinstance(no_val, (int, float)):
            continue  # baris "TOTAL JPL: N titik" dst di ekor tiap sheet
        rec = {"btp_wilayah_kerja": wilayah}
        for col, idx in col_idx.items():
            val = _clean(r[idx]) if idx < len(r) else None
            if col == "no" and val is not None:
                val = int(val)
            if col == "jumlah_kecelakaan" and val is not None:
                val = int(val)
            if col in ("no_jpl", "frekuensi_ka", "headway_ka") and isinstance(val, float):
                # kolom TEXT tapi sumbernya kadang angka murni (mis. NO JPL=98.0,
                # FREKUENSI KA=100.0) -- distring-kan tanpa ".0" palsu.
                val = str(int(val)) if val.is_integer() else str(val)
            rec[col] = val
        out.append(rec)
    return out


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True, read_only=True)
    all_rows = []
    for sheet_name in REGIONAL_SHEETS:
        rows = parse_sheet(wb[sheet_name], sheet_name.title())
        print(f"{sheet_name}: {len(rows)} baris JPL")
        all_rows.extend(rows)

    records = [tuple(row.get(c) for c in TABLE_COLS) for row in all_rows]

    with pg_cursor() as cur:
        cur.execute(Path(__file__).with_name("schema_jpl_prioritas_djka.sql").read_text(encoding="utf-8"))
        cur.execute("DELETE FROM jpl_prioritas_djka")
        col_sql = ", ".join(TABLE_COLS)
        ph_sql = ", ".join(["%s"] * len(TABLE_COLS))
        cur.executemany(f"INSERT INTO jpl_prioritas_djka ({col_sql}) VALUES ({ph_sql})", records)

    print(f"Selesai: {len(records)} baris JPL diimpor ke jpl_prioritas_djka.")


if __name__ == "__main__":
    main()
