"""Analisis Skoring Jalan AWP-1 -- Cost Effectiveness Ratio (CER) per koridor.

Model EKSPERIMENTAL, terpisah dari skor IJD A-E dan NPR. Rumus disalin dari
deck "20261007 PENAMBAHAN PENYEMPURNAAN SIJALAN" slide 10-20 (contekan rumus:
slide 20) dan diverifikasi angka per angka terhadap file sumbernya
`docs/07102026/VfM-koridor-kirim bappenas.xlsx` sheet All (lihat
scripts/import_cer_awp1.py, yang mencetak laporan kecocokan).

Alur (slide 10): kondisi jalan eksisting -> skenario sesudah penanganan ->
indikator eksisting vs sesudah -> manfaat (delta) -> skor 0-10 (dibagi nilai
MAX seluruh koridor) -> TOT SCORE, C SCORE, CER = TOT / C.

Catatan metodologi yang SENGAJA dipertahankan apa adanya dari file sumber
(lihat juga slide 19-22): normalisasi memakai MAX (bukan persentil), MAX
dihitung dari semua koridor termasuk yang Ditolak, dan biaya memakai Biaya
Standar (bukan biaya usulan Pemda).
"""
from typing import Dict, List, Optional

BIAYA_SATUAN_M_PER_KM = {"B": 0.06, "S": 0.2, "RR": 4.0, "RB": 8.0}   # slide 11
DEGRADASI = 0.0139                                                     # slide 13
KECEPATAN_KMJ = {"B": 43.4, "S": 35.7, "RR": 25.5, "RB": 15.3}        # slide 14
BBM_KOEF = (0.0023, -0.3139, 17.499)                                   # BBM = a v^2 + b v + c
FAKTOR_GRK = 2597.86                                                   # GRK = f * BBM / 100
BOK_RP_PER_KM = {"B": 3749, "S": 5075, "RR": 7015, "RB": 9152}        # slide 15
BOBOT_KELAS_JALAN = {"JN": 0.42, "JP": 0.35, "JK": 0.23}              # slide 16
EFEK_KECELAKAAN = 0.30
BOBOT_MANFAAT = {"hpp": 0.50, "bok": 0.15, "pfa": 0.15, "wt": 0.10, "acc": 0.05, "grk": 0.05}  # slide 18


def _kondisi_sesudah(b, s, rr, rb):
    """Slide 13: RR & RB pulih jadi Baik, lalu jalan menurun 1,39%."""
    return {
        "B": (1 - DEGRADASI) * b + rr + rb,
        "S": (1 - DEGRADASI) * s + DEGRADASI * b,
        "RR": DEGRADASI * s,
        "RB": 0.0,
    }


def _rerata(k: Dict[str, float], tabel: Dict[str, float], total: float) -> Optional[float]:
    if not total:
        return None
    return sum(k[x] * tabel[x] for x in ("B", "S", "RR", "RB")) / total


def _bbm(speed):
    a, b, c = BBM_KOEF
    return a * speed ** 2 + b * speed + c


def hitung_komponen(r: dict, harga: Dict[str, float]) -> dict:
    """Komponen mentah satu koridor (sebelum normalisasi). r = input dari sheet All.
    harga = {nama komoditas: Rp/kg} dari sheet price."""
    b, s, rr, rb = (float(r.get(k) or 0) for k in ("b_km", "s_km", "rr_km", "rb_km"))
    total = b + s + rr + rb
    eks = {"B": b, "S": s, "RR": rr, "RB": rb}
    aft = _kondisi_sesudah(b, s, rr, rb)
    biaya_std = sum(eks[k] * BIAYA_SATUAN_M_PER_KM[k] for k in eks)

    pf = []
    for i in (1, 2, 3):
        kom = r.get(f"komoditas_{i}")
        prod = float(r.get(f"produksi_{i}_ton") or 0)
        pf.append(harga.get(kom, 0.0) * prod * 1000 if kom else 0.0)
    hpp = sum(pf)

    out = {"total_km": total, "biaya_std_m": biaya_std, "pf1": pf[0], "pf2": pf[1], "pf3": pf[2], "hpp": hpp,
           "pfa": sum(float(r.get(k) or 0) for k in ("fas_pendidikan", "fas_kesehatan", "fas_pemerintahan", "fas_sppg"))}
    sp_e, sp_a = _rerata(eks, KECEPATAN_KMJ, total), _rerata(aft, KECEPATAN_KMJ, total)
    if sp_e is None:  # koridor tanpa panjang: semua manfaat kondisi = 0
        out.update(speed_eks=None, speed_aft=None, bbm_eks=None, bbm_aft=None, grk_eks=None, grk_aft=None,
                   bok_eks=None, bok_aft=None, wt_eks=None, wt_aft=None, d_grk=0.0, d_bok=0.0, d_wt=0.0)
    else:
        bbm_e, bbm_a = _bbm(sp_e), _bbm(sp_a)
        grk_e, grk_a = FAKTOR_GRK * bbm_e / 100, FAKTOR_GRK * bbm_a / 100
        bok_e, bok_a = _rerata(eks, BOK_RP_PER_KM, total), _rerata(aft, BOK_RP_PER_KM, total)
        wt_e, wt_a = total / sp_e, total / sp_a
        out.update(speed_eks=sp_e, speed_aft=sp_a, bbm_eks=bbm_e, bbm_aft=bbm_a, grk_eks=grk_e, grk_aft=grk_a,
                   bok_eks=bok_e, bok_aft=bok_a, wt_eks=wt_e, wt_aft=wt_a,
                   d_grk=grk_e - grk_a, d_bok=max(bok_e - bok_a, 0.0), d_wt=max(wt_e - wt_a, 0.0))

    jn, jp, jk = (float(r.get(k) or 0) for k in ("jn_km", "jp_km", "jk_km"))
    acc_prov = float(r.get("acc_provinsi") or 0)
    pembagi = BOBOT_KELAS_JALAN["JN"] * jn + BOBOT_KELAS_JALAN["JP"] * jp + BOBOT_KELAS_JALAN["JK"] * jk
    acc_kab = (jk * BOBOT_KELAS_JALAN["JK"]) / pembagi * acc_prov if pembagi else 0.0
    acc_kor = total / jk * acc_kab if jk else 0.0
    out.update(acc_kab=acc_kab, acc_koridor=acc_kor, d_acc=EFEK_KECELAKAAN * acc_kor)
    return out


def hitung_skor(daftar: List[dict]) -> dict:
    """Normalisasi 0-10 (dibagi MAX seluruh koridor) + TOT/C/CER. Mengubah
    `daftar` di tempat (menambah skor_*), mengembalikan nilai MAX yang dipakai."""
    maks = {
        "hpp": max((d["hpp"] for d in daftar), default=0),
        "bok": max((d["d_bok"] for d in daftar), default=0),
        "pfa": max((d["pfa"] for d in daftar), default=0),
        "wt": max((d["d_wt"] for d in daftar), default=0),
        "acc": max((d["d_acc"] for d in daftar), default=0),
        "grk": max((d["d_grk"] for d in daftar), default=0),
        "biaya_std": max((d["biaya_std_m"] for d in daftar), default=0),
    }

    def sk(v, m, guard=False):
        if not m or (guard and not v > 0):
            return 0.0
        return v / m * 10

    for d in daftar:
        d["skor_hpp"] = sk(d["hpp"], maks["hpp"])
        d["skor_bok"] = sk(d["d_bok"], maks["bok"])
        d["skor_pfa"] = sk(d["pfa"], maks["pfa"], guard=True)
        d["skor_wt"] = sk(d["d_wt"], maks["wt"])
        d["skor_acc"] = sk(d["d_acc"], maks["acc"])
        d["skor_grk"] = sk(d["d_grk"], maks["grk"], guard=True)
        d["tot_score"] = sum(BOBOT_MANFAAT[k] * d[f"skor_{k}"] for k in BOBOT_MANFAAT)
        d["c_score"] = d["biaya_std_m"] / maks["biaya_std"] * 10 if maks["biaya_std"] else 0.0
        d["cer_score"] = d["tot_score"] / d["c_score"] if d["c_score"] > 0 else 0.0
    return maks


# --- Preview & Export Skor CER (bahan reviu pemilihan koridor) ---------------
# Pola sama dgn Export Skor IJD: baris per koridor, kolom dikelompokkan, filter
# provinsi (bisa beberapa), peringkat nasional + per provinsi. Membaca tabel
# cer_awp1_koridor apa adanya (skor TIDAK dihitung ulang di sini).

# (grup, label, kolom/alias SQL, jenis) -- jenis menentukan pembulatan preview
# dan number_format xlsx.
EXPORT_KOLOM = [
    ("IDENTITAS", "No.", None, "no"),
    ("IDENTITAS", "ID Koridor (SITIA)", "id_koridor", "id"),
    ("IDENTITAS", "Provinsi", "provinsi", "teks"),
    ("IDENTITAS", "Kabupaten/Kota", "kabupaten_kota", "teks"),
    ("IDENTITAS", "No. Koridor", "no_koridor", "teks"),
    ("IDENTITAS", "Nama Koridor", "nama_koridor", "teks"),
    ("IDENTITAS", "Status Pengajuan", "status_pengajuan", "teks"),
    ("IDENTITAS", "RPJMN", "rpjmn", "teks"),
    ("IDENTITAS", "Tematik", "tematik", "teks"),
    ("HASIL & PERINGKAT", "CER SCORE (TOT ÷ C)", "cer_score", "skor"),
    ("HASIL & PERINGKAT", "Peringkat CER Nasional", "peringkat_cer", "id"),
    ("HASIL & PERINGKAT", "Peringkat CER Provinsi", "peringkat_cer_provinsi", "id"),
    ("HASIL & PERINGKAT", "TOT SCORE (Manfaat, 0-10)", "tot_score", "skor"),
    ("HASIL & PERINGKAT", "Peringkat TOT Nasional", "peringkat_tot", "id"),
    ("HASIL & PERINGKAT", "Peringkat TOT Provinsi", "peringkat_tot_provinsi", "id"),
    ("HASIL & PERINGKAT", "C SCORE (Biaya Relatif, 0-10)", "c_score", "kecil"),
    ("KONDISI & BIAYA", "Panjang Koridor (km)", "panjang_km", "km"),
    ("KONDISI & BIAYA", "Baik (km)", "baik_km", "km"),
    ("KONDISI & BIAYA", "Sedang (km)", "sedang_km", "km"),
    ("KONDISI & BIAYA", "Rusak Ringan (km)", "rusak_ringan_km", "km"),
    ("KONDISI & BIAYA", "Rusak Berat (km)", "rusak_berat_km", "km"),
    ("KONDISI & BIAYA", "Biaya Standar (Rp M)", "biaya_std_m", "rpm"),
    ("KONDISI & BIAYA", "Biaya Usulan Pemda (Rp M, tidak dipakai skor)", "biaya_pemda_m", "rpm"),
    ("MANFAAT (NILAI)", "Komoditas", "komoditas", "teks"),
    ("MANFAAT (NILAI)", "Nilai Produksi / HPP (Rp per tahun)", "hpp_rp", "rp"),
    ("MANFAAT (NILAI)", "Fasilitas Umum Dilewati / PFA (unit)", "pfa", "angka"),
    ("MANFAAT (NILAI)", "Δ BOK (Rp/km)", "d_bok_rp_km", "angka"),
    ("MANFAAT (NILAI)", "Δ Waktu Tempuh (jam)", "d_wt_jam", "skor"),
    ("MANFAAT (NILAI)", "Δ Kecelakaan (kejadian, alokasi estimasi)", "d_acc", "skor"),
    ("MANFAAT (NILAI)", "Δ Emisi GRK (diduga g CO2/km)", "d_grk", "skor"),
    ("MANFAAT (SKOR 0-10)", f"Skor Nilai Produksi (bobot {BOBOT_MANFAAT['hpp']:.0%})", "skor_hpp", "skor"),
    ("MANFAAT (SKOR 0-10)", f"Skor BOK (bobot {BOBOT_MANFAAT['bok']:.0%})", "skor_bok", "skor"),
    ("MANFAAT (SKOR 0-10)", f"Skor Fasilitas Umum (bobot {BOBOT_MANFAAT['pfa']:.0%})", "skor_pfa", "skor"),
    ("MANFAAT (SKOR 0-10)", f"Skor Waktu Tempuh (bobot {BOBOT_MANFAAT['wt']:.0%})", "skor_wt", "skor"),
    ("MANFAAT (SKOR 0-10)", f"Skor Kecelakaan (bobot {BOBOT_MANFAAT['acc']:.0%})", "skor_acc", "skor"),
    ("MANFAAT (SKOR 0-10)", f"Skor Emisi GRK (bobot {BOBOT_MANFAAT['grk']:.0%})", "skor_grk", "skor"),
    ("CATATAN", "Geometri di Peta", "geometri", "teks"),
    ("CATATAN", "Catatan Data", "catatan_data", "teks"),
]
_PEMBULATAN = {"kecil": 5, "skor": 3, "km": 2, "rpm": 2, "rp": 0, "angka": 1}
_FORMAT_XLSX = {"kecil": "0.00000", "skor": "#,##0.000", "km": "#,##0.00", "rpm": "#,##0.00", "rp": "#,##0", "angka": "#,##0.0"}
URUTAN = {"cer": "peringkat_cer", "tot": "peringkat_tot"}


def daftar_provinsi(cur) -> List[dict]:
    cur.execute("SELECT provinsi, count(*) AS jumlah FROM cer_awp1_koridor "
                "WHERE provinsi IS NOT NULL GROUP BY provinsi ORDER BY provinsi")
    return cur.fetchall()


def baris_export(cur, provinsi: List[str], urut: str = "cer") -> List[dict]:
    """Baris tabel CER (dict per koridor), sudah diberi peringkat per provinsi.
    Peringkat provinsi dihitung atas SELURUH koridor provinsi itu sebelum
    difilter, jadi tidak berubah saat provinsi lain ikut dipilih."""
    kunci = URUTAN.get(urut, "peringkat_cer")
    where, params = "", []
    if provinsi:
        where, params = "WHERE provinsi = ANY(%s)", [list(provinsi)]
    cur.execute(f"""
        SELECT * FROM (
            SELECT c.*,
                   CASE WHEN punya_geometri THEN 'Ada' ELSE 'Tidak ada' END AS geometri,
                   row_number() OVER (PARTITION BY provinsi ORDER BY peringkat_cer) AS peringkat_cer_provinsi,
                   row_number() OVER (PARTITION BY provinsi ORDER BY peringkat_tot) AS peringkat_tot_provinsi
            FROM cer_awp1_koridor c
        ) t {where}
        ORDER BY {kunci}""", params)
    return cur.fetchall()


def nilai_baris(r: dict, no: int, bulatkan: bool) -> list:
    out = []
    for _g, _label, kolom, jenis in EXPORT_KOLOM:
        v = no if jenis == "no" else r.get(kolom)
        if bulatkan and isinstance(v, float) and jenis in _PEMBULATAN:
            v = round(v, _PEMBULATAN[jenis])
            if _PEMBULATAN[jenis] == 0:
                v = int(v)
        out.append(v)
    return out


def ringkasan(rows: List[dict], provinsi: List[str], urut: str) -> dict:
    n = len(rows)
    cer = sorted(r["cer_score"] or 0 for r in rows)
    return {
        "cakupan": "Nasional" if not provinsi else ", ".join(provinsi),
        "urut": "CER SCORE (efisiensi)" if urut != "tot" else "TOT SCORE (besarnya manfaat)",
        "jumlah": n,
        "nol_km": sum(1 for r in rows if not ((r["baik_km"] or 0) + (r["sedang_km"] or 0)
                                               + (r["rusak_ringan_km"] or 0) + (r["rusak_berat_km"] or 0))),
        "tanpa_geometri": sum(1 for r in rows if not r["punya_geometri"]),
        "median_cer": cer[n // 2] if n else None,
        "cer_di_atas_50": sum(1 for v in cer if v > 50),
        "diimpor_pada": max((r["diimpor_pada"] for r in rows), default=None),
    }


def _baris_keterangan(info: dict) -> List[tuple]:
    bobot = ", ".join(f"{k.upper()} {v:.0%}" for k, v in BOBOT_MANFAAT.items())
    diimpor = info["diimpor_pada"].strftime("%d-%m-%Y %H:%M") if info["diimpor_pada"] else "-"
    median = f"{info['median_cer']:.2f}".replace(".", ",") if info["median_cer"] is not None else "-"
    return [
        ("Analisis Skoring Jalan AWP-1 — Cost Effectiveness Ratio (CER) per Koridor", None),
        ("Status", "EKSPERIMENTAL — bahan reviu pemilihan koridor. Terpisah dari skor IJD Prioritisasi "
                   "Teknokratik (A–E) dan NPR, bukan keputusan resmi."),
        ("Sumber data", "VfM-koridor-kirim bappenas.xlsx (sheet All + price); rumus deck "
                        "\"20261007 PENAMBAHAN PENYEMPURNAAN SIJALAN\" slide 10–22. Seluruh TOT/C/CER "
                        "dihitung ulang dan cocok 100% dengan nilai di Excel sumber."),
        ("Data diimpor", diimpor),
        ("Cakupan", info["cakupan"]),
        ("Jumlah koridor", info["jumlah"]),
        ("Urutan baris", info["urut"]),
        (None, None),
        ("RUMUS", None),
        ("Biaya Standar", "Baik×0,06 + Sedang×0,2 + Rusak Ringan×4 + Rusak Berat×8 (Rp M per km kondisi)"),
        ("Kondisi sesudah", "Rusak Ringan & Rusak Berat pulih menjadi Baik, lalu jalan menurun 1,39%"),
        ("Nilai Produksi (HPP)", "Σ harga komoditas × produksi (ton) × 1.000, maksimal 3 komoditas"),
        ("Δ BOK, Δ Waktu Tempuh", "Selisih biaya operasi kendaraan (Rp/km) dan waktu tempuh koridor (jam) "
                                 "eksisting − sesudah; negatif dianggap 0"),
        ("Δ Kecelakaan", "Kecelakaan provinsi dialokasikan ke jalan kab/kota (bobot kelas jalan 0,42/0,35/0,23), "
                         "lalu ke koridor menurut panjang; penanganan diasumsikan menurunkan 30%"),
        ("Δ Emisi GRK", "Dari kecepatan → BBM (fungsi kuadrat) → GRK = 2.597,86 × BBM / 100; satuan diduga g CO2/km. Bernilai negatif pada koridor yang sudah Baik seluruhnya (efek degradasi 1,39%), skornya 0"),
        ("Fasilitas Umum (PFA)", "Jumlah fasilitas pendidikan + kesehatan + pemerintahan + SPPG yang dilewati"),
        ("Skor 0–10", "Nilai ÷ nilai terbesar SELURUH koridor × 10"),
        ("TOT SCORE", f"Σ bobot × skor ({bobot})"),
        ("C SCORE", "Biaya Standar ÷ Biaya Standar terbesar × 10"),
        ("CER SCORE", "TOT SCORE ÷ C SCORE (0 bila C SCORE = 0)"),
        (None, None),
        ("CATATAN UNTUK REVIU", None),
        ("TOT vs CER", "Bukan satu ranking tunggal: TOT menunjukkan besarnya manfaat, CER menunjukkan "
                       "efisiensi manfaat per biaya. Kedua urutan bisa sangat berbeda (deck slide 19)."),
        ("CER ekstrem", f"Koridor pendek berbiaya kecil bisa mendapat CER sangat tinggi. Dalam cakupan ini: "
                        f"median CER {median}, {info['cer_di_atas_50']} koridor dengan CER > 50."),
        ("Normalisasi MAX", "Pembagi skor = nilai terbesar dari seluruh koridor (termasuk yang Ditolak), sehingga "
                            "satu koridor ekstrem menekan skor koridor lain. Sengaja dipertahankan sesuai Excel "
                            "sumber sampai pemilik model memutuskan (deck slide 22)."),
        ("Panjang kondisi 0 km", f"{info['nol_km']} koridor dalam cakupan ini: manfaat kondisi jalan = 0, "
                                 "menunggu konfirmasi data sumber."),
        ("Tanpa geometri", f"{info['tanpa_geometri']} koridor dalam cakupan ini tidak ada di layer PETA KORIDOR "
                           "(tidak bisa ditampilkan di peta)."),
        ("Biaya usulan Pemda", "Ditampilkan sebagai pembanding saja; skor memakai Biaya Standar."),
        ("Peringkat provinsi", "Dihitung terhadap seluruh koridor di provinsi tersebut."),
    ]


def tulis_xlsx(rows: List[dict], info: dict) -> bytes:
    import io

    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Skor CER Koridor"
    tebal = Font(bold=True, size=10)
    tengah = Alignment(horizontal="center", vertical="center", wrap_text=True)
    abu = PatternFill("solid", fgColor="F2F2F2")

    # Baris 1: grup kolom (merge), baris 2: judul kolom.
    awal = 1
    for i in range(1, len(EXPORT_KOLOM) + 1):
        if i == len(EXPORT_KOLOM) or EXPORT_KOLOM[i][0] != EXPORT_KOLOM[i - 1][0]:
            ws.cell(row=1, column=awal, value=EXPORT_KOLOM[awal - 1][0])
            if i > awal:
                ws.merge_cells(start_row=1, start_column=awal, end_row=1, end_column=i)
            awal = i + 1
    for c, (_g, label, _k, _j) in enumerate(EXPORT_KOLOM, start=1):
        ws.cell(row=2, column=c, value=label)
    for baris in (1, 2):
        for c in range(1, len(EXPORT_KOLOM) + 1):
            sel = ws.cell(row=baris, column=c)
            sel.font, sel.alignment, sel.fill = tebal, tengah, abu
    ws.row_dimensions[2].height = 45

    for no, r in enumerate(rows, start=1):
        ws.append(nilai_baris(r, no, bulatkan=False))
    for c, (_g, label, _k, jenis) in enumerate(EXPORT_KOLOM, start=1):
        huruf = get_column_letter(c)
        lebar = {"no": 6, "id": 12, "teks": 18}.get(jenis, 15)
        if label in ("Nama Koridor", "Catatan Data", "Komoditas"):
            lebar = 36
        ws.column_dimensions[huruf].width = lebar
        if jenis in _FORMAT_XLSX:
            for sel in ws.iter_rows(min_row=3, min_col=c, max_col=c):
                sel[0].number_format = _FORMAT_XLSX[jenis]
    ws.freeze_panes = "C3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(EXPORT_KOLOM))}{len(rows) + 2}"

    ket = wb.create_sheet("Keterangan")
    for a, b in _baris_keterangan(info):
        ket.append([a, b])
        if a and b is None:
            ket.cell(row=ket.max_row, column=1).font = Font(bold=True, size=11 if ket.max_row == 1 else 10)
        else:
            ket.cell(row=ket.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
            ket.cell(row=ket.max_row, column=1).alignment = Alignment(vertical="top")
    ket.column_dimensions["A"].width = 24
    ket.column_dimensions["B"].width = 110

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
