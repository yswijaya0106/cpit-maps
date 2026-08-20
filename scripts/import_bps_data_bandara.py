# -*- coding: utf-8 -*-
"""Impor atribut detail 251 bandara nasional (docs/New/3. UDARA/
Attributes_ Data Bandara.xlsx, sheet "Sheet1") ke tabel referensi
bps_data_bandara -- lihat scripts/schema_bps_data_bandara.sql untuk
skema & alasan tabel terpisah (bukan menimpa layer BANDARA existing),
dan docs/kajian_data_baru_docs_new.md §5 untuk hasil telaah datanya.
TIDAK terkait usulan_inpres/IJD.

Sumbernya punya baris kelanjutan tanpa nomor urut (`NO` kosong) untuk
bandara dengan >1 taxiway -- dideteksi lewat kolom Taxiway (idx 11)
terisi sementara kolom lain kosong, digabung ke baris utama bandara
terakhir yang punya `NO`. Baris separator murni (semua kolom kosong/
whitespace non-breaking-space) dilewati.

Koordinat "Titik Koordinat" berformat DMS notasi Indonesia (LU/LS/BT/
BB) -- diparse ke lat/lon desimal, disimpan di samping teks aslinya.

Idempotent: DELETE + INSERT ulang seluruh tabel tiap run.

Usage (venv aktif):
    python scripts/import_bps_data_bandara.py
"""
import io
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

import openpyxl

from db import db_cursor as pg_cursor  # noqa: E402

REPO_ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = (
    REPO_ROOT / "docs" / "New" / "3. UDARA-20260820T015257Z-1-001"
    / "3. UDARA" / "Attributes_ Data Bandara.xlsx"
)
SCHEMA_PATH = Path(__file__).resolve().parent / "schema_bps_data_bandara.sql"

# Kolom (0-based tuple index, diverifikasi manual thd baris data --
# header 2-baris merge di sumber xlsx tidak bisa dipercaya index-nya
# begitu saja): 0=NO, 1=Nama Bandara, 2=Hirarki, 3=Kelas, 4=Provinsi,
# 5=Kabupaten, 6=Status, 7=Operator, 8=Runway Length, 9=Runway Width,
# 10=Apron Area, 11=Taxiway, 12=Terminal Penumpang, 13=Demand Pax,
# 14=Terminal Kargo, 15=Critical Aircraft, 16=Kapasitas Eksisting Valid,
# 17=Kapasitas Eksisting Estimasi, 18=Titik Koordinat, 19=KP, 20=KD.
COL_NO, COL_NAMA, COL_HIRARKI, COL_KELAS = 0, 1, 2, 3
COL_PROVINSI, COL_KABUPATEN, COL_STATUS, COL_OPERATOR = 4, 5, 6, 7
COL_RW_LEN, COL_RW_WID, COL_APRON, COL_TAXIWAY = 8, 9, 10, 11
COL_TERM_PAX, COL_DEMAND, COL_TERM_KARGO, COL_CRITICAL_AC = 12, 13, 14, 15
COL_KAP_VALID, COL_KAP_ESTIMASI = 16, 17
COL_KOORDINAT, COL_KP, COL_KD = 18, 19, 20

_DMS_RE = re.compile(r"(\d+)\D+(\d+)\D+([\d.,]+)\D*(N|S|E|W|LU|LS|BT|BB)", re.IGNORECASE)
_NEGATIVE_HEMI = {"S", "W", "LS", "BB"}


def _parse_dms(s):
    if not s:
        return None
    m = _DMS_RE.search(str(s))
    if not m:
        return None
    deg, minute, sec, hemi = m.groups()
    sec = sec.replace(",", ".")
    try:
        val = float(deg) + float(minute) / 60 + float(sec) / 3600
    except ValueError:
        return None
    return -val if hemi.upper() in _NEGATIVE_HEMI else val


def _num(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None  # mis. "Tidak Terdefinisi"


def _text(v):
    if v is None:
        return None
    s = str(v).replace("\xa0", "").strip()
    return s or None


def _load_rows():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True, read_only=True)
    ws = wb["Sheet1"]
    rows = []
    current = None  # dict baris bandara aktif (menampung baris kelanjutan taxiway)

    for row in ws.iter_rows(min_row=4, values_only=True):
        no = row[COL_NO]
        if isinstance(no, (int, float)):
            if current is not None:
                rows.append(current)
            lat = _parse_dms(row[COL_KOORDINAT])
            lon = None
            if row[COL_KOORDINAT]:
                # DMS dua pasang (lat lalu lon) dlm satu string -- cari pasangan kedua
                m_all = list(_DMS_RE.finditer(str(row[COL_KOORDINAT])))
                if len(m_all) >= 2:
                    d, mi, se, hemi = m_all[1].groups()
                    se = se.replace(",", ".")
                    try:
                        lon = float(d) + float(mi) / 60 + float(se) / 3600
                        if hemi.upper() in _NEGATIVE_HEMI:
                            lon = -lon
                    except ValueError:
                        lon = None
            current = {
                "no": int(no), "nama_bandara": _text(row[COL_NAMA]),
                "hirarki": _text(row[COL_HIRARKI]), "kelas": _text(row[COL_KELAS]),
                "provinsi": _text(row[COL_PROVINSI]), "kabupaten": _text(row[COL_KABUPATEN]),
                "status": _text(row[COL_STATUS]), "operator": _text(row[COL_OPERATOR]),
                "runway_length_m": _num(row[COL_RW_LEN]), "runway_width_m": _num(row[COL_RW_WID]),
                "apron_area_m2": _num(row[COL_APRON]),
                "taxiway": [t for t in [_text(row[COL_TAXIWAY])] if t],
                "terminal_penumpang_m2": _num(row[COL_TERM_PAX]), "demand_pax": _num(row[COL_DEMAND]),
                "terminal_kargo": _text(row[COL_TERM_KARGO]), "critical_aircraft": _text(row[COL_CRITICAL_AC]),
                "kapasitas_eksisting_valid": _num(row[COL_KAP_VALID]),
                "kapasitas_eksisting_estimasi": _num(row[COL_KAP_ESTIMASI]),
                "titik_koordinat_dms": _text(row[COL_KOORDINAT]),
                "lat": lat, "lon": lon,
                "kode_provinsi": int(row[COL_KP]) if isinstance(row[COL_KP], (int, float)) else None,
                "kode_kabupaten": int(row[COL_KD]) if isinstance(row[COL_KD], (int, float)) else None,
            }
        elif current is not None:
            tw = _text(row[COL_TAXIWAY])
            if tw:
                current["taxiway"].append(tw)

    if current is not None:
        rows.append(current)

    for r in rows:
        r["taxiway"] = "; ".join(r["taxiway"]) or None
    return rows


def main():
    if not XLSX_PATH.exists():
        print(f"GAGAL: tidak ditemukan {XLSX_PATH}")
        sys.exit(1)

    with pg_cursor() as cur:
        cur.execute(SCHEMA_PATH.read_text(encoding="utf-8"))

    print(f"Membaca {XLSX_PATH.name}...")
    rows = _load_rows()
    print(f"  {len(rows)} bandara")

    cols = ["no", "nama_bandara", "hirarki", "kelas", "provinsi", "kabupaten", "status",
            "operator", "runway_length_m", "runway_width_m", "apron_area_m2", "taxiway",
            "terminal_penumpang_m2", "demand_pax", "terminal_kargo", "critical_aircraft",
            "kapasitas_eksisting_valid", "kapasitas_eksisting_estimasi",
            "titik_koordinat_dms", "lat", "lon", "kode_provinsi", "kode_kabupaten"]
    placeholders = ", ".join(["%s"] * len(cols))

    with pg_cursor() as cur:
        cur.execute("DELETE FROM bps_data_bandara")
        cur.executemany(
            f"INSERT INTO bps_data_bandara ({', '.join(cols)}) VALUES ({placeholders})",
            [tuple(r[c] for c in cols) for r in rows],
        )

    print("\nSelesai: bps_data_bandara di-refresh.")


if __name__ == "__main__":
    main()
