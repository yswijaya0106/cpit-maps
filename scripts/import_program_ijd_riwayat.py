"""Import Riwayat Program IJD 2023-2026 (daftar kegiatan yang DIPROGRAMKAN/DPP final,
bukan usulan) -> tabel program_ijd_riwayat.

Sumber: docs/03102026/Data Revisi R1_Riwayat_IJD_2023-2026_Gabungan.xlsx (sheet
"RIWAYAT IJD 2023-2026", 2.083 kegiatan, Rp48,8 T) -- file yang sama dengan
yang dipakai deck Bappenas "20261002 Preparation, Implementation and Validation
of IJD" (angka biaya/km per provinsi di deck tereproduksi persis dari file ini).

Kenapa tabel sendiri (tidak digabung ke usulan_inpres_riwayat / dpp_ijd_2025):
- usulan_inpres_riwayat = USULAN SITIA (alokasi tiap tahap usulan); nilai_dpp
  di sana untuk 2023/2024 tidak sama dengan daftar final ini (2023: 1.123 baris
  33 T di SITIA vs 512 kegiatan 14,6 T di sini).
- dpp_ijd_2025 hanya satu tahun (646 kegiatan); 2025 di file ini 534 kegiatan,
  semuanya cocok alokasi persis dgn dpp_ijd_2025 -- 112 sisanya belum jelas.

Normalisasi (kolom asli tetap disimpan):
- kode wilayah: "KD Kab/Kota" sumber kadang kode provinsi (xx00) utk kegiatan
  kab, Banten sebagian 3600, Papua Barat Daya 98xx (master BPS 92xx). Kode
  kabupaten dicari ulang dari NAMA via wilayah_cocok.PencocokKabupaten; kode
  sumber dipakai hanya bila nama tak tercocokkan tapi kodenya ada di master.
  Baris "Provinsi X" (pengusul provinsi) -> kode_kabupaten NULL.
- kategori: Jenis Kegiatan PENINGKATAN -> 'preservasi', PEMBANGUNAN ->
  'pembangunan' (sama dgn deck: nama "Peningkatan" berganti "Preservasi" 2024).
- tahap: huruf besar, spasi dirapikan (2026 kosong di sumber -> NULL).
- usulan_inpres_id (hanya 2026): alokasi = usulan_inpres.alokasi_usulan_kompetensi
  persis DAN provinsi sama, dan hanya bila kandidatnya tunggal.
- catatan_data: anomali yang terlihat (alokasi 0, biaya/km tak wajar, kode
  dikoreksi) -- baris TIDAK dibuang, supaya total tetap sama dgn sumber.

DELETE + INSERT seluruh tabel (file = daftar lengkap), aman di-rerun.

Usage (venv aktif, .env berisi PG_*):
    python scripts/import_program_ijd_riwayat.py [path.xlsx]
"""
import os
import re
import sys
from pathlib import Path

import openpyxl
import psycopg
from psycopg.rows import dict_row

sys.path.insert(0, str(Path(__file__).resolve().parent))
from wilayah_cocok import PencocokKabupaten, kunci_provinsi  # noqa: E402

BASE_DIR = Path(__file__).resolve().parent.parent
DEFAULT_XLSX = BASE_DIR / "docs" / "03102026" / "Data Revisi R1_Riwayat_IJD_2023-2026_Gabungan.xlsx"

HEADER = {
    "No": "no_sumber", "Provinsi": "provinsi", "Tahun": "tahun", "Tahap": "tahap",
    "Nama Kegiatan": "nama_kegiatan", "Jenis Kegiatan": "jenis_kegiatan",
    "Status Jalan": "status_jalan", "Kab/Kota": "kab_kota", "KD Kab/Kota": "kd_kab_kota_sumber",
    "Panjang Jalan (km)": "panjang_jalan_km", "Panjang Jembatan (m)": "panjang_jembatan_m",
    "Alokasi (Rp)": "alokasi_rp", "Tematik": "tematik",
}
KATEGORI = {"PENINGKATAN": "preservasi", "PEMBANGUNAN": "pembangunan"}

DDL = """
CREATE TABLE IF NOT EXISTS program_ijd_riwayat (
    id SERIAL PRIMARY KEY,
    no_sumber INTEGER,
    tahun SMALLINT NOT NULL,
    tahap TEXT,
    provinsi TEXT,
    nama_kegiatan TEXT,
    jenis_kegiatan TEXT,
    kategori TEXT,                 -- 'preservasi' | 'pembangunan'
    status_jalan TEXT,             -- K / P apa adanya
    kab_kota TEXT,
    kd_kab_kota_sumber INTEGER,    -- kode di file sumber (tidak selalu kode BPS)
    kode_provinsi SMALLINT,        -- ID BPS hasil normalisasi
    kode_kabupaten INTEGER,        -- ID BPS; NULL utk kegiatan provinsi
    dasar_kode TEXT,               -- 'nama' | 'kode_sumber' | 'provinsi' | NULL
    panjang_jalan_km NUMERIC,
    panjang_jembatan_m NUMERIC,
    alokasi_rp NUMERIC,
    tematik TEXT,
    usulan_inpres_id BIGINT,       -- 2026 saja, lihat docstring script
    catatan_data TEXT,
    sumber_file TEXT,
    diimpor_at TIMESTAMPTZ DEFAULT now()
);
CREATE INDEX IF NOT EXISTS idx_program_ijd_riwayat_tahun ON program_ijd_riwayat (tahun);
CREATE INDEX IF NOT EXISTS idx_program_ijd_riwayat_kab ON program_ijd_riwayat (kode_kabupaten);
"""


def connect():
    from dotenv import load_dotenv
    load_dotenv(BASE_DIR / ".env")
    return psycopg.connect(
        host=os.environ.get("PG_HOST", "127.0.0.1"), port=int(os.environ.get("PG_PORT", "5432")),
        user=os.environ.get("PG_USER", "postgres"), password=os.environ.get("PG_PASS", ""),
        dbname=os.environ.get("PG_DB", "route_gis"), row_factory=dict_row,
    )


def baca_xlsx(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["RIWAYAT IJD 2023-2026"] if "RIWAYAT IJD 2023-2026" in wb.sheetnames else wb.active
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    hilang = [h for h in HEADER if h not in header]
    if hilang:
        sys.exit(f"Kolom tidak ditemukan di {path}: {hilang}")
    idx = {HEADER[h]: header.index(h) for h in HEADER}
    out = []
    for r in rows:
        if r is None or all(v is None for v in r):
            continue
        rec = {k: r[i] if i < len(r) else None for k, i in idx.items()}
        if rec["tahun"] is None:
            continue
        out.append(rec)
    return out


def _teks(v):
    if v is None:
        return None
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s or None


def _angka(v):
    try:
        return float(v) if v is not None and str(v).strip() != "" else None
    except (TypeError, ValueError):
        return None


def normalisasi(rows, pencocok, kab_master, prov_master):
    kode_kab_master = {int(m["kode_kabupaten"]) for m in kab_master}
    for r in rows:
        catatan = []
        r["tahun"] = int(r["tahun"])
        r["no_sumber"] = int(r["no_sumber"]) if r["no_sumber"] is not None else None
        for k in ("provinsi", "nama_kegiatan", "jenis_kegiatan", "status_jalan", "kab_kota", "tematik"):
            r[k] = _teks(r[k])
        r["tahap"] = _teks(r["tahap"]).upper() if _teks(r["tahap"]) else None
        r["jenis_kegiatan"] = (r["jenis_kegiatan"] or "").upper() or None
        r["kategori"] = KATEGORI.get(r["jenis_kegiatan"])
        for k in ("panjang_jalan_km", "panjang_jembatan_m", "alokasi_rp"):
            r[k] = _angka(r[k])
        kd = int(r["kd_kab_kota_sumber"]) if _angka(r["kd_kab_kota_sumber"]) is not None else None
        r["kd_kab_kota_sumber"] = kd

        r["kode_kabupaten"], r["kode_provinsi"], r["dasar_kode"] = None, None, None
        nama = r["kab_kota"] or ""
        if re.match(r"^\s*prov(insi|\.)?\s", nama, re.I):
            kp = prov_master.get(kunci_provinsi(nama)) or prov_master.get(kunci_provinsi(r["provinsi"]))
            if kp:
                r["kode_provinsi"], r["dasar_kode"] = kp, "provinsi"
        else:
            m = pencocok.cari(r["provinsi"], nama)
            if m:
                r["kode_kabupaten"], r["dasar_kode"] = int(m["kode_kabupaten"]), "nama"
            elif kd in kode_kab_master:
                r["kode_kabupaten"], r["dasar_kode"] = kd, "kode_sumber"
            if r["kode_kabupaten"]:
                r["kode_provinsi"] = r["kode_kabupaten"] // 100
        if r["kode_kabupaten"] and kd and kd != r["kode_kabupaten"]:
            catatan.append(f"kode sumber {kd} dikoreksi -> {r['kode_kabupaten']}")
        if not r["kode_provinsi"]:
            catatan.append("wilayah tidak tercocokkan")

        km, alok = r["panjang_jalan_km"] or 0, r["alokasi_rp"] or 0
        if alok <= 0:
            catatan.append("alokasi 0")
        elif km > 0:
            per_km = alok / 1e9 / km
            if per_km < 0.5:
                catatan.append(f"biaya/km sangat rendah ({per_km:.3f} Rp M/km, cek satuan)")
            elif per_km > 100:
                catatan.append(f"biaya/km sangat tinggi ({per_km:.0f} Rp M/km, ruas pendek)")
        if km <= 0 and (r["panjang_jembatan_m"] or 0) <= 0:
            catatan.append("tanpa panjang jalan & jembatan")
        r["catatan_data"] = "; ".join(catatan) or None
    return rows


def tautkan_usulan_2026(cur, rows):
    """Program 2026 -> usulan_inpres lewat alokasi Kompetensi yang persis sama di provinsi
    yang sama. Banyak nilai alokasi kembar (angka bulat, mis. Rp20 M) -> kandidat
    disaring bertahap: kab/kota sama, lalu panjang penanganan Kompetensi sama.
    Tetap tidak tunggal -> tidak ditautkan (lebih baik kosong daripada salah)."""
    cur.execute("SELECT id, kode_provinsi, kode_kabupaten, alokasi_usulan_kompetensi AS a, "
                "panjang_penanganan_kompetensi AS pjg FROM usulan_inpres WHERE alokasi_usulan_kompetensi > 0")
    idx = {}
    for u in cur.fetchall():
        idx.setdefault((round(float(u["a"])), u["kode_provinsi"]), []).append(u)
    n = 0
    for r in rows:
        r["usulan_inpres_id"] = None
        if r["tahun"] != 2026 or not r["alokasi_rp"]:
            continue
        cand = idx.get((round(r["alokasi_rp"]), r["kode_provinsi"])) or []
        if len(cand) > 1 and r["kode_kabupaten"]:
            cand = [u for u in cand if u["kode_kabupaten"] == r["kode_kabupaten"]] or cand
        if len(cand) > 1 and r["panjang_jalan_km"]:
            cand = [u for u in cand if u["pjg"] is not None
                    and abs(float(u["pjg"]) - r["panjang_jalan_km"]) < 0.005] or cand
        if len(cand) == 1:
            r["usulan_inpres_id"] = cand[0]["id"]
            n += 1
    return n


def main():
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    if not path.exists():
        sys.exit(f"File tidak ada: {path}")
    rows = baca_xlsx(path)
    conn = connect()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT DISTINCT kode_provinsi, provinsi, kode_kabupaten, kabupaten_kota FROM penduduk_kecamatan")
            kab_master = cur.fetchall()
            prov_master = {kunci_provinsi(m["provinsi"]): int(m["kode_provinsi"]) for m in kab_master}
            normalisasi(rows, PencocokKabupaten(kab_master), kab_master, prov_master)
            n_tautan = tautkan_usulan_2026(cur, rows)
            cur.execute(DDL)
            cur.execute("DELETE FROM program_ijd_riwayat")
            kolom = ["no_sumber", "tahun", "tahap", "provinsi", "nama_kegiatan", "jenis_kegiatan", "kategori",
                     "status_jalan", "kab_kota", "kd_kab_kota_sumber", "kode_provinsi", "kode_kabupaten",
                     "dasar_kode", "panjang_jalan_km", "panjang_jembatan_m", "alokasi_rp", "tematik",
                     "usulan_inpres_id", "catatan_data"]
            cur.executemany(
                f"INSERT INTO program_ijd_riwayat ({', '.join(kolom)}, sumber_file) "
                f"VALUES ({', '.join(['%s'] * (len(kolom) + 1))})",
                [tuple(r[k] for k in kolom) + (path.name,) for r in rows])
        conn.commit()
        with conn.cursor() as cur:
            cur.execute("SELECT tahun, COUNT(*) n, ROUND(SUM(alokasi_rp)/1e12, 2) t, "
                        "COUNT(*) FILTER (WHERE kode_provinsi IS NULL) tanpa_wil, "
                        "COUNT(*) FILTER (WHERE catatan_data IS NOT NULL) bercatatan "
                        "FROM program_ijd_riwayat GROUP BY 1 ORDER BY 1")
            for r in cur.fetchall():
                print(f"{r['tahun']}: {r['n']} kegiatan, Rp {r['t']} T, tanpa wilayah {r['tanpa_wil']}, "
                      f"bercatatan {r['bercatatan']}")
            cur.execute("SELECT dasar_kode, COUNT(*) n FROM program_ijd_riwayat GROUP BY 1 ORDER BY 2 DESC")
            print("Dasar kode wilayah:", {r["dasar_kode"]: r["n"] for r in cur.fetchall()})
        print(f"Program 2026 tertaut ke usulan_inpres: {n_tautan}")
    finally:
        conn.close()


if __name__ == "__main__":
    main()
