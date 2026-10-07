# -*- coding: utf-8 -*-
"""Impor & hitung Analisis Skoring Jalan AWP-1 (CER per koridor).

Sumber: docs/07102026/VfM-koridor-kirim bappenas.xlsx (sheet All + price),
deck "20261007 PENAMBAHAN PENYEMPURNAAN SIJALAN" slide 5-22. Rumus ada di
cer_awp1.py (root repo), BUKAN disalin dari nilai Excel: skrip ini menghitung
ulang dari input, lalu MEMVALIDASI hasilnya terhadap nilai TOT/C/CER yang
tersimpan di Excel. Bila ada yang tidak cocok, skrip berhenti tanpa menulis
(kecuali --paksa).

Menulis:
1. tabel cer_awp1_koridor (schema_cer_awp1.sql), DELETE + INSERT;
2. layer overlay "Koridor AWP-1" di map_layers, bucket KORIDOR AWP-1, satu layer
   per provinsi. Geometri = ruas layer PETA KORIDOR dgn ID_KORIDOR yg sama
   (ID SITIA, unik per koridor per kab/kota). Garis pink via _warna (deck slide 3).

Usage (venv aktif, PG_* di .env):
    python scripts/import_cer_awp1.py [--xlsx PATH] [--paksa]
Restart server (atau tunggu cache layer) setelah rerun.
"""
import argparse
import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl  # noqa: E402
from dotenv import load_dotenv  # noqa: E402
from psycopg.types.json import Json  # noqa: E402

load_dotenv(REPO_ROOT / ".env")
import cer_awp1  # noqa: E402
from db import db_cursor  # noqa: E402
from wilayah_cocok import PencocokKabupaten  # noqa: E402

XLSX_BAWAAN = REPO_ROOT / "docs" / "07102026" / "VfM-koridor-kirim bappenas.xlsx"
SCHEMA = Path(__file__).resolve().parent / "schema_cer_awp1.sql"
BUCKET = "KORIDOR AWP-1"
LAYER = "Koridor AWP-1"
WARNA = "#ec4899"  # pink (deck slide 3: "Koridor AWP-1")

# Indeks kolom sheet All (baris 1 = header), diverifikasi 7 Okt 2026.
K = dict(id=1, provinsi=2, kab=3, no_kor=4, nama=5, rpjmn=6, tematik=7, status=10,
         panjang=28, b=29, s=30, rr=31, rb=32, biaya_pemda=33,
         kom1=38, prod1=40, kom2=42, prod2=44, kom3=46, prod3=48,
         fas_pend=57, fas_kes=58, fas_pem=59, fas_sppg=60,
         acc=81, jn=82, jp=83, jk=84,
         x_tot=129, x_c=131, x_cer=133)


class HargaKomoditas(dict):
    """VLOOKUP Excel tidak peka huruf besar/kecil."""

    def get(self, k, d=None):
        return super().get(str(k).strip().casefold(), d) if k else d


def baca_excel(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    harga = HargaKomoditas()
    for r in wb["price"].iter_rows(min_row=1, values_only=True):
        if r[1] and isinstance(r[3], (int, float)):
            harga.setdefault(str(r[1]).strip().casefold(), float(r[3]))
    rows = []
    for r in wb["All"].iter_rows(min_row=2, values_only=True):
        if not isinstance(r[0], (int, float)) or r[K["id"]] is None:
            continue
        rows.append({
            "id_koridor": int(r[K["id"]]), "provinsi": r[K["provinsi"]], "kabupaten_kota": r[K["kab"]],
            "no_koridor": r[K["no_kor"]], "nama_koridor": r[K["nama"]], "status_pengajuan": r[K["status"]],
            "rpjmn": r[K["rpjmn"]], "tematik": r[K["tematik"]],
            "panjang_km": r[K["panjang"]], "b_km": r[K["b"]], "s_km": r[K["s"]], "rr_km": r[K["rr"]],
            "rb_km": r[K["rb"]], "biaya_pemda_m": r[K["biaya_pemda"]],
            "komoditas_1": r[K["kom1"]], "produksi_1_ton": r[K["prod1"]],
            "komoditas_2": r[K["kom2"]], "produksi_2_ton": r[K["prod2"]],
            "komoditas_3": r[K["kom3"]], "produksi_3_ton": r[K["prod3"]],
            "fas_pendidikan": r[K["fas_pend"]], "fas_kesehatan": r[K["fas_kes"]],
            "fas_pemerintahan": r[K["fas_pem"]], "fas_sppg": r[K["fas_sppg"]],
            "acc_provinsi": r[K["acc"]], "jn_km": r[K["jn"]], "jp_km": r[K["jp"]], "jk_km": r[K["jk"]],
            "_excel": {"tot_score": r[K["x_tot"]], "c_score": r[K["x_c"]], "cer_score": r[K["x_cer"]]},
        })
    return rows, harga


def validasi(rows, hasil):
    beda = []
    for r, h in zip(rows, hasil):
        for k, ev in r["_excel"].items():
            if not isinstance(ev, (int, float)) or abs(h[k] - ev) > max(1e-6, 1e-6 * abs(ev)):
                beda.append((r["id_koridor"], k, ev, h[k]))
    return beda


def peringkat(hasil, kunci):
    urut = sorted(range(len(hasil)), key=lambda i: -hasil[i][kunci])
    out = [0] * len(hasil)
    for pos, i in enumerate(urut, start=1):
        out[i] = pos
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", type=Path, default=XLSX_BAWAAN)
    ap.add_argument("--paksa", action="store_true", help="tetap tulis walau hasil tidak cocok dgn Excel")
    args = ap.parse_args()

    rows, harga = baca_excel(args.xlsx)
    hasil = [cer_awp1.hitung_komponen(r, harga) for r in rows]
    maks = cer_awp1.hitung_skor(hasil)
    beda = validasi(rows, hasil)
    print(f"{len(rows)} koridor dibaca. Pembagi MAX: " + ", ".join(f"{k}={v:,.4g}" for k, v in maks.items()))
    print(f"Validasi TOT/C/CER vs nilai Excel: {3 * len(rows) - len(beda)} cocok, {len(beda)} beda")
    for b in beda[:10]:
        print("  BEDA", b)
    if beda and not args.paksa:
        sys.exit("Berhenti: hasil hitung ulang tidak sama dgn Excel (pakai --paksa utk tetap menulis).")
    rk_tot, rk_cer = peringkat(hasil, "tot_score"), peringkat(hasil, "cer_score")

    with db_cursor() as cur:
        cur.execute(SCHEMA.read_text(encoding="utf-8"))
        cur.execute("SELECT DISTINCT kode_provinsi, provinsi, kode_kabupaten, kabupaten_kota FROM penduduk_kecamatan")
        pencocok = PencocokKabupaten(cur.fetchall())
        cur.execute("""SELECT (attrs->>'ID_KORIDOR')::numeric::int AS id, count(*) AS n
                       FROM map_layers WHERE layer = 'PETA KORIDOR' AND attrs->>'ID_KORIDOR' ~ '^[0-9.]+$'
                       GROUP BY 1""")
        ruas = {g["id"]: g["n"] for g in cur.fetchall()}

        baris, tak_cocok_wilayah = [], 0
        for i, (r, h) in enumerate(zip(rows, hasil)):
            m = pencocok.cari(r["provinsi"], r["kabupaten_kota"])
            tak_cocok_wilayah += m is None
            catatan = []
            if not h["total_km"]:
                catatan.append("panjang kondisi 0 km: manfaat kondisi jalan = 0")
            if r["id_koridor"] not in ruas:
                catatan.append("tidak ada geometri di layer PETA KORIDOR")
            komoditas = ", ".join(str(r[f"komoditas_{j}"]) for j in (1, 2, 3) if r[f"komoditas_{j}"])
            baris.append((
                r["id_koridor"], r["provinsi"], r["kabupaten_kota"],
                int(m["kode_provinsi"]) if m else None, int(m["kode_kabupaten"]) if m else None,
                r["no_koridor"], r["nama_koridor"], r["status_pengajuan"], r["rpjmn"], r["tematik"],
                r["panjang_km"], r["b_km"], r["s_km"], r["rr_km"], r["rb_km"], r["biaya_pemda_m"], h["biaya_std_m"],
                komoditas or None, h["hpp"], h["pfa"], h["d_bok"], h["d_wt"], h["d_acc"], h["d_grk"],
                h["skor_hpp"], h["skor_bok"], h["skor_pfa"], h["skor_wt"], h["skor_acc"], h["skor_grk"],
                h["tot_score"], h["c_score"], h["cer_score"], rk_tot[i], rk_cer[i],
                r["id_koridor"] in ruas, "; ".join(catatan) or None,
            ))
        cur.execute("DELETE FROM cer_awp1_koridor")
        cur.executemany(
            "INSERT INTO cer_awp1_koridor (id_koridor, provinsi, kabupaten_kota, kode_provinsi, kode_kabupaten, "
            "no_koridor, nama_koridor, status_pengajuan, rpjmn, tematik, panjang_km, baik_km, sedang_km, "
            "rusak_ringan_km, rusak_berat_km, biaya_pemda_m, biaya_std_m, komoditas, hpp_rp, pfa, d_bok_rp_km, "
            "d_wt_jam, d_acc, d_grk, skor_hpp, skor_bok, skor_pfa, skor_wt, skor_acc, skor_grk, tot_score, "
            "c_score, cer_score, peringkat_tot, peringkat_cer, punya_geometri, catatan_data) VALUES ("
            + ", ".join(["%s"] * 37) + ")", baris)

        # Layer peta: ruas PETA KORIDOR + atribut CER koridornya, satu layer per provinsi.
        # Geometri disalin di SQL (JOIN ke tabel sementara atribut), tidak lewat Python.
        cur.execute("DELETE FROM map_layers WHERE provinsi = %s", (BUCKET,))
        cur.execute("DELETE FROM map_layer_meta WHERE provinsi = %s", (BUCKET,))
        cur.execute("CREATE TEMP TABLE tmp_awp1 (id_koridor INTEGER PRIMARY KEY, attrs JSONB) ON COMMIT DROP")
        cur.executemany("INSERT INTO tmp_awp1 VALUES (%s, %s)", [(r["id_koridor"], Json({
            "No. Koridor": r["no_koridor"], "Nama Koridor": r["nama_koridor"],
            "Kab/Kota": r["kabupaten_kota"], "Status Pengajuan": r["status_pengajuan"],
            "CER (manfaat per biaya)": round(h["cer_score"], 2),
            "Peringkat CER nasional": f"{rk_cer[i]} dari {len(rows)}",
            "TOT SCORE (manfaat)": round(h["tot_score"], 3),
            "Peringkat manfaat nasional": f"{rk_tot[i]} dari {len(rows)}",
            "C SCORE (biaya relatif)": round(h["c_score"], 3),
            "Biaya Std (Rp M)": round(h["biaya_std_m"], 2),
            "Panjang koridor (km)": round(h["total_km"], 2),
            "Komoditas": ", ".join(str(r[f"komoditas_{j}"]) for j in (1, 2, 3) if r[f"komoditas_{j}"]) or "-",
            "Catatan": "Model eksperimental AWP-1 (deck 20261007), terpisah dari skor IJD & NPR",
            "_warna": WARNA, "_lebar": 3,
        })) for i, (r, h) in enumerate(zip(rows, hasil)) if r["id_koridor"] in ruas])
        cur.execute("""INSERT INTO map_layers (provinsi, kabupaten, layer, attrs, geom)
                       SELECT %s, replace(initcap(p.provinsi), 'Di ', 'DI '), %s,
                              t.attrs || jsonb_build_object('Nama Ruas', p.attrs->>'NAMA_RUAS'), p.geom
                       FROM map_layers p
                       JOIN tmp_awp1 t ON t.id_koridor = (p.attrs->>'ID_KORIDOR')::numeric::int
                       WHERE p.layer = 'PETA KORIDOR' AND p.attrs->>'ID_KORIDOR' ~ '^[0-9.]+$'""",
                    (BUCKET, LAYER))
        cur.execute("""INSERT INTO map_layer_meta (provinsi, kabupaten, layer, label, feature_count, size_mb, source_shp)
                       SELECT provinsi, kabupaten, layer, layer, count(*),
                              round((sum(pg_column_size(attrs) + ST_MemSize(geom)) / 1048576.0)::numeric, 2), %s
                       FROM map_layers WHERE provinsi = %s GROUP BY provinsi, kabupaten, layer""",
                    (str(args.xlsx.relative_to(REPO_ROOT)) + " + layer PETA KORIDOR", BUCKET))
        cur.execute("SELECT count(*) AS n, count(DISTINCT kabupaten) AS p FROM map_layers WHERE provinsi = %s", (BUCKET,))
        lapis = cur.fetchone()

    n_geo = sum(1 for b in baris if b[-2])
    print(f"Ditulis: {len(baris)} koridor ke cer_awp1_koridor ({n_geo} bergeometri, "
          f"{tak_cocok_wilayah} kab/kota tak terpetakan ke kode BPS)")
    print(f"Layer peta '{LAYER}': {lapis['n']} ruas di {lapis['p']} provinsi")


if __name__ == "__main__":
    main()
